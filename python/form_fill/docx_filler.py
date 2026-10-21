# -*- coding: utf-8 -*-
"""docx 底层写入器（移植自 form_fill_v2，原样移植仅换日志框架）。

按 schema target 路径定位格/嵌套表，提供各 fill_type 的写入原语：
路径定位与 template_analyzer 的 target 语义完全一致（物理 tc 索引），
Schema 即接口（docs/2.4forms-design.md §2.1）。
"""
import copy
import logging
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell

logger = logging.getLogger(__name__)

FONT_NAME = '仿宋'
# 「中文标签+冒号+空白」——与 template_analyzer._BLANK_FIELD 保持一致（blank 编号口径统一）
_BLANK_FIELD = re.compile(r'([\u4e00-\u9fa5]{2,10})[:：][\s_]{2,}([\u4e00-\u9fa5]{1,6})?')
_RC = re.compile(r'r(\d+)c(\d+)')


# ============================== 路径定位 ==============================

def _cell(tc_elm) -> _Cell:
    """tc 元素包装为 python-docx _Cell（段落/runs 操作不依赖 parent）"""
    return _Cell(tc_elm, None)


def _tc_at(tbl_elm, r: int, c: int):
    return tbl_elm.findall(qn('w:tr'))[r].findall(qn('w:tc'))[c]


def resolve_node(doc: Document, target: str):
    """解析 schema target → (node, blank_n)
    - 'top0/r4c1'          → (tc元素, None)
    - 'top0/r3c1/nested1'  → (tbl元素, None)（table_rows 的表路径）
    - 'top0/r4c1#blank1'   → (tc元素, 1)（选项内嵌填空）
    """
    main, _, suffix = target.partition('#')
    blank_n = int(suffix[len('blank'):]) if suffix else None
    parts = main.split('/')
    tbl = doc.tables[int(parts[0][3:])]._tbl
    node = tbl
    for i in range(1, len(parts)):
        seg = parts[i]
        if seg.startswith('nested'):
            r, c = (int(x) for x in _RC.fullmatch(parts[i - 1]).groups())
            tbl = _tc_at(tbl, r, c).findall(qn('w:tbl'))[int(seg[6:])]
            node = tbl
        else:
            r, c = (int(x) for x in _RC.fullmatch(seg).groups())
            node = _tc_at(tbl, r, c)
    return node, blank_n


# ============================== 段内区间替换（跨 run 定位） ==============================

def _replace_span(para, start: int, end: int, value: str):
    """把段内拼接文本 [start, end) 区间替换为 value（定位到具体 run，start==end 时为插入）"""
    off = 0
    done = False
    for run in para.runs:
        r_start, r_end = off, off + len(run.text)
        off = r_end
        if r_end <= start or r_start >= end:
            continue
        s = max(start - r_start, 0)
        e = min(end - r_start, len(run.text))
        if not done:
            run.text = run.text[:s] + value + run.text[e:]
            done = True
        else:
            run.text = run.text[:s] + run.text[e:]
    if not done and para.runs:   # 边界：区间在全部 run 之后 → 追加
        para.runs[-1].text += value


def _copy_run_format(src_run, dst_run):
    """复制 run 格式（字体/字号/加粗，含中文字体名）"""
    if src_run.font.name:
        dst_run.font.name = src_run.font.name
        rpr = dst_run._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = rpr.makeelement(qn('w:rFonts'), {})
            rpr.insert(0, rfonts)
        rfonts.set(qn('w:eastAsia'), src_run.font.name)
    if src_run.font.size:
        dst_run.font.size = src_run.font.size
    if src_run.font.bold is not None:
        dst_run.font.bold = src_run.font.bold


# ============================== 写入原语 ==============================

def write_text(tc_elm, value: str):
    """清格写入（text 长标题填写区 / doc_section 长文），\n 分段，保留首 run 字体"""
    cell = _cell(tc_elm)
    first_font = next((p.runs[0] for p in cell.paragraphs if p.runs), None)

    for p in list(cell.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    first_para = cell.paragraphs[0]
    for run in list(first_para.runs):
        run._element.getparent().remove(run._element)

    for i, line in enumerate(value.split('\n') if value else ['']):
        para = first_para if i == 0 else cell.add_paragraph()
        run = para.add_run(line)
        if first_font is not None:
            _copy_run_format(first_font, run)
        else:
            run.font.name = FONT_NAME
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.find(qn('w:rFonts'))
            if rfonts is None:
                rfonts = rpr.makeelement(qn('w:rFonts'), {})
                rpr.insert(0, rfonts)
            rfonts.set(qn('w:eastAsia'), FONT_NAME)


def _replace_check_symbol(cell, option: str) -> bool:
    """单选项勾选替换 □→☑，支持 □ 与选项文字跨 run 的 3 种分布、引号样式归一"""
    if not option:
        return False
    norm = option.replace('\u201c', '"').replace('\u201d', '"')

    for para in cell.paragraphs:
        runs = para.runs
        for run in runs:   # 情况1：□与选项同 run
            text = run.text
            norm_text = text.replace('\u201c', '"').replace('\u201d', '"')
            for target in (f'□{norm}', f'□ {norm}'):
                idx = norm_text.find(target)
                if idx != -1:
                    run.text = text[:idx] + '☑' + text[idx + 1:]
                    return True
        for i in range(len(runs) - 1):   # 情况2：□在 run 末尾、选项在下个 run 开头
            if runs[i].text.rstrip().endswith('□'):
                remaining = runs[i].text.rstrip()[:-1]
                norm_next = runs[i + 1].text.lstrip().replace('\u201c', '"').replace('\u201d', '"')
                if norm_next.startswith(norm):
                    runs[i].text = remaining + '☑'
                    return True
        for i, run in enumerate(runs):   # 情况3：□与选项前缀同 run、选项跨 run
            text = run.text
            idx = text.find('□')
            while idx != -1:
                after = text[idx + 1:].lstrip()
                norm_after = after.replace('\u201c', '"').replace('\u201d', '"')
                if norm_after and norm.startswith(norm_after):
                    combined = ''.join(r.text for r in runs[i + 1:])
                    combined = combined.replace('\u201c', '"').replace('\u201d', '"')
                    if combined.startswith(norm[len(norm_after):]):
                        run.text = text[:idx] + '☑' + text[idx + 1:]
                        return True
                idx = text.find('□', idx + 1)
    return False


def check_options(tc_elm, selected: list) -> list:
    """choice/checkbox：逐项勾选替换，支持多选；返回未命中选项（写入填表报告）"""
    cell = _cell(tc_elm)
    missed = []
    for opt in selected or []:
        if not _replace_check_symbol(cell, str(opt)):
            missed.append(opt)
    if missed:
        logger.warning('[DocxFiller] 勾选未命中: %s', missed)
    return missed


def fill_after_label(tc_elm, label: str, value: str) -> bool:
    """label_group 锚点插值：定位含 label 主体（冒号前部分，anchor 可能带"标签：提示"全串）
    的段落，把「标签：」后的内容整段替换为 value（「标签：提示文案」→「标签：值」）"""
    cell = _cell(tc_elm)
    loc = re.escape(label.split('：')[0].split(':')[0])
    for para in cell.paragraphs:
        full = ''.join(r.text for r in para.runs)
        m = re.search(loc, full)
        if not m:
            continue
        colon = full.find('：', m.end())
        if colon == -1:
            colon = full.find(':', m.end())
        if colon == -1:
            continue
        _replace_span(para, colon + 1, len(full), value)
        return True
    logger.warning('[DocxFiller] 锚点未命中: %s', label)
    return False


def fill_blank(tc_elm, n: int, value: str, label: str = None) -> bool:
    """#blank{n}：格内「中文标签+冒号+空白」位写入数值（排除"其他"位，口径与 analyzer 一致）。
    优先按标签主体（field label 去括号/单位后的中文部分）匹配，否则取第 n 个空白位
    （序号在先行 blank 已写入后会平移，仅作兜底）；单位文字保留，只写数值"""
    cell = _cell(tc_elm)
    loc = None
    if label:
        m = re.match(r'([\u4e00-\u9fa5]+)', label)
        loc = m.group(1) if m else None
    matches = []
    for para in cell.paragraphs:
        full = ''.join(r.text for r in para.runs)
        for m in _BLANK_FIELD.finditer(full):
            if m.group(1) == '其他':
                continue
            matches.append((para, m))
    chosen = None
    if loc:
        chosen = next(((p, m) for p, m in matches if m.group(1) == loc), None)
    if chosen is None and len(matches) >= n:
        chosen = matches[n - 1]
    if chosen is None:
        logger.warning('[DocxFiller] blank#%d(%s) 定位失败', n, label)
        return False
    para, m = chosen
    blank_start = m.end(1) + 1
    blank_end = m.start(2) if m.group(2) else m.end()
    _replace_span(para, blank_start, blank_end, value)
    return True


def fill_table(tbl_elm, rows: list, header_rows: int = 1):
    """table_rows：数据行从 header_rows 开始逐行逐格写入；预置行不足时 deepcopy
    末数据行追加（多余预置空行保留，是否清理后置）"""
    trs = tbl_elm.findall(qn('w:tr'))
    data_trs = trs[header_rows:]
    if not data_trs:
        logger.warning('[DocxFiller] 数据表无数据行模板')
        return
    while len(data_trs) < len(rows):
        new_tr = copy.deepcopy(data_trs[-1])
        data_trs[-1].addnext(new_tr)
        data_trs.append(new_tr)
    for tr, vals in zip(data_trs, rows):
        for tc, v in zip(tr.findall(qn('w:tc')), vals):
            if v not in (None, ''):
                write_text(tc, str(v))


def apply_field(doc: Document, field: dict, value) -> dict:
    """按 fill_type 分发写入。返回写入结果（missed/ok 供填表报告）。
    value 约定（来自批次 LLM 生成 / 人工修订，经 _normalize_value 归一）：
    - text: str（blank 带 unit 时只给数值）
    - choice/checkbox: {'selected': [..], 'other': str|None}（other 走 fill_after_label）
    - label_group: {sub_key: value}（按 sub_fields anchor 定位）
    - table_rows: [[列值...]]
    - doc_section: str（多段 \n 分隔）"""
    node, blank_n = resolve_node(doc, field['target'])
    ftype = field['fill_type']
    result = {'field_key': field['field_key'], 'ok': True, 'missed': []}

    if field['target'].startswith('top') and '#blank' in field['target']:
        result['ok'] = fill_blank(node, blank_n, str(value), label=field.get('label'))
    elif ftype in ('choice', 'checkbox'):
        payload = value if isinstance(value, dict) else {'selected': [value]}
        result['missed'] = check_options(node, payload.get('selected') or [])
        other = payload.get('other')
        if other:
            # 优先"其他："锚；无则降级用已勾选项的冒号（如"异常情况说明："）
            if not fill_after_label(node, '其他', str(other)):
                for opt in (payload.get('selected') or []):
                    if fill_after_label(node, str(opt), str(other)):
                        break
                else:
                    result['missed'].append('other')
    elif ftype == 'label_group':
        anchors = {sf['sub_key']: sf['anchor'] for sf in field.get('sub_fields', [])}
        for sub_key, sub_val in (value or {}).items():
            if sub_val in (None, ''):
                continue
            if not fill_after_label(node, anchors.get(sub_key, sub_key), str(sub_val)):
                result['missed'].append(sub_key)
    elif ftype == 'table_rows':
        fill_table(node, value or [], header_rows=field.get('header_rows', 1))
    else:   # text / doc_section
        if value not in (None, ''):
            write_text(node, str(value))
    return result
