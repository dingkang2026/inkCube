# -*- coding: utf-8 -*-
"""模板分析器（移植自 form_fill_v2，LLM 调用换 LLMClient）。

阶段一：OOXML 结构提取（纯规则，8 级决策树分类，docs/2.4forms-design.md §2.1）
阶段二：LLM 语义映射 + 批次规划（prompt 见 form_fill/prompts.py）
校验兜底 + Schema 组装（§3.3）

定位模型：XML 路径 top{i}/r{ri}c{ci}/nested{j}/r{ri}c{ci}（tbl/r/tc 均为同层 findall 顺序索引）
"""
import json
import logging
import re

from docx import Document
from docx.oxml.ns import qn

from llm.client import LLMClient
from form_fill.prompts import TEMPLATE_ANALYSIS_PROMPT, build_analysis_user_prompt

logger = logging.getLogger(__name__)

MAX_NESTED_DEPTH = 3                     # 嵌套超过 3 层整表 unmapped（§2.4）
SHORT_LABEL_MAX = 12                     # 短标签阈值
LABEL_MAX = 40                           # 左标签上限（覆盖"长标题+空白填写区"模式）
HEADER_CELL_MAX = 15                     # 表头单元格文本阈值
DOC_MIN_LEN = 20                         # doc_section 最短文本
INSTRUCTION_VERBS = ('说明', '介绍', '列示', '分析', '总结', '扼要', '展示', '列举', '阐述')
CHECK_MARKS = ('□', '☐', '○')
SYMBOL_FONTS = ('wingdings', 'webdings', 'symbol')   # 符号字体：勾选符无法可靠解析
SECTION_TITLE_L1 = re.compile(r'^[一二三四五六七八九十]+、')
SECTION_TITLE_L2 = re.compile(r'^（[一二三四五六七八九十]+）')
KEY_PATTERN = re.compile(r'^[a-z][a-z0-9_]*$')
VALID_FILL_TYPES = ('text', 'label_group', 'table_rows', 'doc_section', 'choice', 'checkbox')
LOCKED_FILL_TYPES = ('choice', 'checkbox', 'table_rows', 'label_group')  # 结构类型锁定，LLM 不得改判


class InvalidTemplateError(Exception):
    """格式不符（status=invalid）：.doc 老格式损坏、无规则表格等"""


# ============================== 阶段一：结构提取 ==============================

def _tc_info(tc) -> dict:
    """采集单个 tc 的分类所需信息"""
    paras = [''.join(t.text or '' for t in p.findall('.//' + qn('w:t'))).strip()
             for p in tc.findall(qn('w:p'))]
    fonts = set()
    for rpr in tc.findall('.//' + qn('w:rPr')):
        for f in rpr.findall(qn('w:rFonts')):
            for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
                v = f.get(qn(attr))
                if v:
                    fonts.add(v.lower())
    tcpr = tc.find(qn('w:tcPr'))
    vmerge = tcpr.find(qn('w:vMerge')) if tcpr is not None else None
    # <w:vMerge/> 无 val 属性时 OOXML 默认为 continue（续行）
    vmerge_cont = vmerge is not None and vmerge.get(qn('w:val')) != 'restart'
    grid_span = 1
    g = tcpr.find(qn('w:gridSpan')) if tcpr is not None else None
    if g is not None and (g.get(qn('w:val')) or '').isdigit():
        grid_span = int(g.get(qn('w:val')))
    return {
        'paras': [p for p in paras if p],
        'text': ' '.join(p for p in paras if p),
        'fonts': fonts,
        'vmerge_cont': vmerge_cont,
        'grid_span': grid_span,
        'nested_tbls': tc.findall(qn('w:tbl')),
    }


def _placeholder_of(text: str):
    """占位样式识别：underline / bracket / None"""
    if '___' in text or '__' in text:
        return 'underline'
    if re.search(r'【[^】]*】', text):
        return 'bracket'
    return None


def _is_short_label(text: str) -> bool:
    return 0 < len(text.strip()) <= SHORT_LABEL_MAX


def _check_count(text: str) -> int:
    return sum(text.count(m) for m in CHECK_MARKS)


def _is_label_like(text: str) -> bool:
    """可作左标签：非占位/非勾选/无指令动词且 ≤LABEL_MAX
    （放宽长标题，使其右侧空白填写区能成字段，如"公司财务状况（近三年收入、研发、利润情况）"）"""
    t = text.strip()
    if not t or len(t) > LABEL_MAX:
        return False
    return (not _placeholder_of(t) and _check_count(t) == 0
            and not any(v in t for v in INSTRUCTION_VERBS))


_BLANK_FIELD = re.compile(r'([\u4e00-\u9fa5]{2,10})[:：][\s_]{2,}([\u4e00-\u9fa5]{1,6})?')


def _extract_option(o: str):
    """解析单个选项 → (options, blanks)：
    - "标签：___单位"内嵌填空拆出 blanks（如 融资金额：___万元 → 独立填空字段），
      选项名保留填空标签（如 "租赁：___平方米" → "租赁"，使勾选替换能命中模板原文）
    - "其他[：___]"还原为独立选项（修复与前一选项粘连）"""
    blanks = [{'label': m.group(1), 'unit': m.group(2) or ''}
              for m in _BLANK_FIELD.finditer(o) if m.group(1) != '其他']
    # 仅当片段以「标签：填空」开头时，标签才作选项名（"租赁：___平方米"→"租赁"，
    # 使勾选替换能命中模板原文）；片段中部的填空行（如"融资金额：___万元 ；投前估值：___亿元"）
    # 只是填空不进 options
    core = o.lstrip()
    head = _BLANK_FIELD.match(core)
    if blanks and head:
        text = head.group(1)
    else:
        text = _BLANK_FIELD.sub('', core)
    text = re.sub(r'[\s_]+', ' ', text).strip(' ；;，, ：:')
    if '其他' in text:
        parts = [p.strip(' ；;，, ：:') for p in text.split('其他')]
        options = [p for p in parts if p] + ['其他']
    else:
        options = [text] if text else []
    return options, blanks


def _split_instruction_sections(instruction: str) -> list:
    """按"一、二、…"中文序号切分写作指引为子节列表（预解析入 schema，填表按子节生成）；
    不足 2 段或存在过短段时不切（交由填表侧自行处理）"""
    parts = [p.strip() for p in re.split(r'(?=[一二三四五六七八九十]+、)', instruction) if p.strip()]
    if len(parts) < 2 or any(len(p) < 4 for p in parts):
        return []
    return parts


class _TemplateExtractor:
    """阶段一执行器：递归遍历 → 决策树分类 → 候选清单"""

    def __init__(self):
        self.candidates = []
        self.unmapped = []
        self.top_tables = 0
        self.nested_tables = 0

    # ---------- 主入口 ----------
    def extract(self, doc: Document):
        if not doc.tables:
            raise InvalidTemplateError('文档中无任何规则表格')
        self.top_tables = len(doc.tables)
        for i, tbl in enumerate(doc.tables):
            self._walk_table(tbl._tbl, f'top{i}', depth=0, base_chain=[])
        return self

    # ---------- 表级遍历 ----------
    def _walk_table(self, tbl_elm, path: str, depth: int, base_chain: list):
        if depth > MAX_NESTED_DEPTH:
            self.unmapped.append({'target': path, 'reason': f'嵌套超过{MAX_NESTED_DEPTH}层'})
            return
        rows = []
        for tr in tbl_elm.findall(qn('w:tr')):
            rows.append([_tc_info(tc) for tc in tr.findall(qn('w:tc'))])

        # 表头探测：整表判定为数据表 → 产出 1 个 table_rows 字段
        detected = self._detect_data_table(rows)
        if detected:
            header_rows, columns = detected
            if depth > 0:
                self.nested_tables += 1
            self.candidates.append({
                'target': path,
                'structural_type': 'table_rows',
                'placeholder': None,
                'context': {
                    'left_label': None,
                    'section_chain': list(base_chain),
                    'table_header': columns,
                    'instruction': None,
                },
                'options': None,
                'preset_rows': max(len(rows) - header_rows, 0),
                '_data_path': f'{path}/r{{{header_rows}..{len(rows) - 1}}}',
                '_header_rows': header_rows,
            })
            return
        if depth > 0:
            self.nested_tables += 1

        # 逐行逐 tc 分类（8 级决策树）
        for ri, row in enumerate(rows):
            left_label = None
            for ci, info in enumerate(row):
                if info['vmerge_cont']:
                    continue  # vMerge 续行格归属 restart 格
                self._classify_tc(info, f'{path}/r{ri}c{ci}', depth, base_chain, left_label)
                if info['text'] and _is_label_like(info['text']):
                    left_label = info['text']

    # ---------- 表头探测 ----------
    def _detect_data_table(self, rows):
        """首行(或前两行)为表头 且 数据区存在数据行 → 数据表，返回 (header_rows, columns)
        - 表头行判定：非空格 ≥2 且均为短文本、无占位/勾选/纯数字（允许空格：vMerge续行/装饰列）
        - 数据行判定：全空行，或非空格=1且为纯数字（预置序号行）——
          标签值两列表（每行1个非数字标签）不会命中，避免误判（如基本信息表）
        - vMerge 续行格不参与判定（非独立逻辑格）"""
        def header_like(row):
            non_empty = [c['text'] for c in row if c['text'] and not c['vmerge_cont']]
            if len(non_empty) < 2:
                return False
            return all(len(t) <= HEADER_CELL_MAX
                       and not _placeholder_of(t) and _check_count(t) == 0
                       and not re.fullmatch(r'[\d.,%\s]+', t)   # 纯数字格视为数据（预置序号）
                       for t in non_empty)

        def is_data_row(row):
            non_empty = [c['text'] for c in row
                         if c['text'] and not c['vmerge_cont']]
            if not non_empty:
                return True
            return (len(non_empty) == 1
                    and bool(re.fullmatch(r'[\d.,%\s]+', non_empty[0])))

        if not rows or not header_like(rows[0]):
            return None
        header_rows = 1
        # 两级表头：第二行也是表头样式，且其后存在数据行
        if len(rows) > 2 and header_like(rows[1]) and any(is_data_row(r) for r in rows[2:]):
            header_rows = 2
        elif not any(is_data_row(r) for r in rows[1:]):
            return None
        columns = self._expand_columns(rows[0], rows[1] if header_rows == 2 else None)
        return header_rows, columns

    def _expand_columns(self, first_row, second_row):
        """按逻辑列展开表头：gridSpan 展开、vMerge续行继承、二级表头对位合并，丢弃纯空装饰列"""
        col_names = []
        for c in first_row:
            col_names.extend([c['text']] * c['grid_span'])
        width = len(col_names)
        if second_row:
            second_names = []
            for c in second_row:
                second_names.extend(['' if (c['vmerge_cont'] or not c['text']) else c['text']] * c['grid_span'])
            second_names = (second_names + [''] * width)[:width]
            col_names = [f"{a}/{b}" if (b and b != a) else a
                         for a, b in zip(col_names, second_names)]
        return [name for name in col_names if name]

    # ---------- 单元格分类（8 级决策树，命中即止） ----------
    def _classify_tc(self, info, target: str, depth: int, base_chain: list, left_label):
        text, paras, fonts = info['text'], info['paras'], info['fonts']

        # 规则1：container（含嵌套表格）
        if info['nested_tbls']:
            chain, rest_text = self._split_sections(paras, base_chain)
            for j, ntbl in enumerate(info['nested_tbls']):
                self._walk_table(ntbl, f'{target}/nested{j}', depth + 1, chain)
            # 格内剩余段落若含写作指引 → 额外产出 1 个 doc_section
            if len(rest_text) >= DOC_MIN_LEN and any(v in rest_text for v in INSTRUCTION_VERBS):
                self._add_doc_section(target, chain, rest_text)
            return

        # Wingdings 类符号字体 → 勾选符无法可靠解析（§2.4）
        if fonts & set(SYMBOL_FONTS):
            self.unmapped.append({'target': target, 'reason': '符号字体(Wingdings)勾选符'})
            return

        # 规则2：勾选符（选项拆分：内嵌填空成独立 blank 候选、"其他"粘连还原；
        # 单勾选符 checkbox 同样拆分——覆盖"勾选+填数"复合形态，如"租赁：___平方米"）
        checks = _check_count(text)
        if checks >= 1:
            ftype = 'choice' if checks >= 2 else 'checkbox'
            options, blanks = [], []
            for o in [x.strip() for x in re.split(r'[□☐○]', text) if x.strip()]:
                opts, blks = _extract_option(o)
                options.extend(opts)
                blanks.extend(blks)
            self._append(target, ftype, text, base_chain, left_label, options=options)
            for n, b in enumerate(blanks, 1):
                bc = self._append(f'{target}#blank{n}', 'text', '', base_chain, left_label,
                                  placeholder='underline')
                bc['_blank_label'] = b['label']
                if b['unit']:
                    bc['_blank_unit'] = b['unit']
            return

        # 规则3：label_group（多段落均为"短标签："模式）
        colon_paras = [p for p in paras if p.rstrip().endswith(('：', ':')) and len(p) <= 30]
        if len(colon_paras) >= 2:
            # 以"（"开头的括注是提示尾巴（如"（包括但不限于以上内容）"），不产锚
            anchors = [p.rstrip().rstrip('：:').strip() for p in paras
                       if p.strip() and not p.strip().startswith('（')]
            cand = self._append(target, 'label_group', text, base_chain, left_label)
            cand['_anchors'] = [a for a in anchors if a]
            return

        # 规则4：下划线/【】占位
        ph = _placeholder_of(text)
        if ph:
            self._append(target, 'text', text, base_chain, left_label, placeholder=ph)
            return

        # 规则5：空格 → 上下文判定
        if not text.strip():
            if left_label and _is_label_like(left_label):
                self._append(target, 'text', text, base_chain, left_label)
            else:
                self.unmapped.append({'target': target, 'reason': '空格且无左标签', 'context': {'self_text': ''}})
            return

        # 规则6：短文本无占位 → label（仅作上下文）
        if _is_short_label(text):
            return

        # 规则7：长文本含指令动词 → doc_section
        if len(text) >= DOC_MIN_LEN and any(v in text for v in INSTRUCTION_VERBS):
            self._add_doc_section(target, base_chain, text)
            return

        # 规则8：兜底 label，交给阶段二二次判断
        self._append(target, 'label', text, base_chain, left_label)

    # ---------- 工具 ----------
    def _split_sections(self, paras: list, base_chain: list):
        """容器格段落切分：标题入章节链，其余聚合为写作指引"""
        chain = list(base_chain)
        sub_chain: list = []
        instruction_parts = []
        for p in paras:
            if SECTION_TITLE_L1.match(p) and len(p) <= 40:
                sub_chain = [p]
            elif SECTION_TITLE_L2.match(p) and len(p) <= 40:
                sub_chain = (sub_chain[:1] if sub_chain else []) + [p]
            else:
                instruction_parts.append(p)
        chain += sub_chain
        return chain, '\n'.join(instruction_parts)

    def _add_doc_section(self, target, chain, instruction):
        self._append(target, 'doc_section', instruction[:60], chain, None, instruction=instruction)

    def _append(self, target, structural_type, text, chain, left_label, *,
                placeholder=None, options=None, instruction=None) -> dict:
        cand = {
            'target': target,
            'structural_type': structural_type,
            'placeholder': placeholder,
            'context': {
                'left_label': left_label,
                'section_chain': list(chain),
                'table_header': None,
                'instruction': instruction,
            },
            'options': options,
            'preset_rows': None,
        }
        if text:
            cand['_self_text'] = text[:60]
        self.candidates.append(cand)
        return cand


# ============================== 阶段二：语义映射 ==============================

def _parse_llm_json(content: str) -> dict:
    """解析 LLM 返回的严格 JSON（容忍 ```json 围栏）"""
    text = content.strip()
    fence = re.search(r'```(?:json)?\s*(.+?)```', text, re.S)
    if fence:
        text = fence.group(1).strip()
    start = text.find('{')
    end = text.rfind('}')
    if start < 0 or end <= start:
        raise ValueError('LLM 输出中未找到 JSON 对象')
    return json.loads(text[start:end + 1])


async def _semantic_map(template_name: str, candidates: list,
                        llm_client: LLMClient,
                        prev_mapping: dict = None, user_hint: str = None) -> dict:
    """LLM 一次调用：字段语义化 + 批次规划；
    重解析时附带上次映射基线与用户提示做局部修正（修订模式）"""
    user_prompt = build_analysis_user_prompt(template_name, candidates,
                                             prev_mapping=prev_mapping, user_hint=user_hint)
    messages = [
        {'role': 'system', 'content': TEMPLATE_ANALYSIS_PROMPT},
        {'role': 'user', 'content': user_prompt},
    ]
    result = await llm_client.chat_full(messages)
    usage = result.get('usage', {})
    raw = result.get('content') or ''
    if not raw.strip():
        raise RuntimeError(f'LLM 语义映射返回空响应, usage={usage}')
    logger.info('[TemplateAnalyzer] 语义映射完成, usage=%s, length=%d', usage, len(raw))
    return _parse_llm_json(raw)


# ============================== 校验与组装 ==============================

def _default_batches(fields: list) -> list:
    """默认规则分批（兜底）：按文档顺序，每 12 字段或遇 doc_section 边界成批"""
    batches, cur = [], []

    def flush():
        if not cur:
            return
        name = next((f['context']['section_chain'][-1] for f in cur
                     if f['context']['section_chain']), cur[0].get('label') or '未分区')
        batches.append({
            'seq': len(batches) + 1,
            'name': name,
            'field_keys': [f['field_key'] for f in cur],
            'depends_on': [],
        })
        cur.clear()

    for f in fields:
        if f['fill_type'] == 'doc_section' and cur:
            flush()
        cur.append(f)
        if len(cur) >= 12:
            flush()
    if len(cur) == 1 and batches:   # 尾部单字段小批并入上一批
        batches[-1]['field_keys'].append(cur[0]['field_key'])
        cur.clear()
    flush()
    return batches


def _batches_valid(batches: list, all_keys: set) -> bool:
    """批次校验：覆盖性 / 引用存在 / 无环 / 体量"""
    if not batches:
        return False
    seqs = [b.get('seq') for b in batches]
    covered = [k for b in batches for k in b.get('field_keys', [])]
    if len(covered) != len(set(covered)) or set(covered) != all_keys:
        return False
    if any(len(b.get('field_keys', [])) > 15 for b in batches):
        return False
    by_seq = {b.get('seq'): b for b in batches}
    state = {}

    def dfs(seq):   # 白灰黑三色检环
        if state.get(seq) == 1:
            return False
        if state.get(seq) == 2:
            return True
        state[seq] = 1
        for dep in by_seq.get(seq, {}).get('depends_on', []) or []:
            if dep not in by_seq or not dfs(dep):
                return False
        state[seq] = 2
        return True

    return all(dfs(s) for s in seqs)


def _assemble(template_name: str, extractor: _TemplateExtractor, llm: dict) -> dict:
    """合并结构候选与 LLM 语义输出 → 落盘 Schema（§3.3）"""
    indexed = list(enumerate(extractor.candidates))
    by_index = {i: c for i, c in indexed}
    skipped = set(llm.get('skipped_indexes') or [])
    mapped_indexes = set()
    fields = []
    used_keys = set()

    for f in llm.get('fields', []):
        idx = f.get('index')
        cand = by_index.get(idx)
        if not cand:
            continue
        mapped_indexes.add(idx)
        if idx in skipped:
            extractor.unmapped.append({'target': cand['target'], 'reason': 'LLM判定为装饰格'})
            continue
        # field_key 规范化
        key = f.get('field_key') or ''
        if not KEY_PATTERN.match(key) or key in used_keys:
            key = f'field_{idx}'
            n = 2
            while key in used_keys:
                key = f'field_{idx}_{n}'
                n += 1
        used_keys.add(key)
        if cand['structural_type'] in LOCKED_FILL_TYPES:
            fill_type = cand['structural_type']
        else:
            fill_type = f.get('fill_type') if f.get('fill_type') in VALID_FILL_TYPES else cand['structural_type']
        fill_mode = 'rule_fill' if f.get('fill_mode') == 'rule_fill' else 'kb_fill'
        if cand.get('_blank_label'):   # 选项内嵌填空：用填空名作 label
            label = cand['_blank_label'] + (f"({cand.get('_blank_unit')})"
                                            if cand.get('_blank_unit') else '')
        else:
            label = (cand['context'].get('left_label')
                     or cand.get('_self_text')
                     or (cand['context'].get('table_header') or [''])[0]
                     or key)
        field = {
            'field_key': key,
            'label': label,
            'fill_type': fill_type,
            'fill_mode': fill_mode,
            'target': cand['target'],
            'placeholder': cand['placeholder'],
            'context': cand['context'],
            'kb_query': (f.get('kb_query') or [])[:3],
            'length_hint': f.get('length_hint'),
            'confidence': f.get('confidence') or 'low',
        }
        if fill_mode == 'rule_fill' and f.get('rule_hint'):
            field['rule_hint'] = f['rule_hint']
        if fill_type in ('choice', 'checkbox') and isinstance(f.get('multi_select'), bool):
            field['multi_select'] = f['multi_select']   # 单选/多选语义
        if cand.get('_blank_unit'):
            field['unit'] = cand['_blank_unit']   # 填空单位：填充只写数值不重复单位
        if cand.get('options'):
            field['options'] = cand['options']
        if fill_type == 'label_group':
            # 以阶段一 anchors 为准逐锚点落 sub_fields，LLM 输出按 anchor 匹配补充 kb_query；
            # 无 anchors 时才沿用 LLM 的 sub_fields
            anchors = cand.get('_anchors') or []
            llm_subs = {s.get('anchor'): s
                        for s in (f.get('sub_fields') or []) if s.get('anchor')}
            field['sub_fields'] = [
                {'sub_key': (llm_subs[a].get('sub_key') or f'sub_{n + 1}') if a in llm_subs
                            else f'sub_{n + 1}',
                 'anchor': a,
                 'kb_query': (llm_subs[a].get('kb_query') or [])[:3] if a in llm_subs else []}
                for n, a in enumerate(anchors)
            ] or [
                {'sub_key': s.get('sub_key') or f'sub_{n + 1}',
                 'anchor': s.get('anchor') or '',
                 'kb_query': (s.get('kb_query') or [])[:3]}
                for n, s in enumerate(f.get('sub_fields') or [])
            ]
        if fill_type == 'table_rows':
            field['columns'] = cand['context'].get('table_header') or []
            field['preset_rows'] = cand.get('preset_rows')
            field['data_path'] = cand.get('_data_path')
            field['header_rows'] = cand.get('_header_rows', 1)
        if fill_type == 'doc_section':
            field['instruction'] = instr = cand['context'].get('instruction') or ''
            secs = _split_instruction_sections(instr)   # 子节预解析
            if secs:
                field['sections'] = secs
        fields.append(field)

    # LLM 未映射的非 label 候选 → unmapped（不阻塞落盘）
    for i, cand in indexed:
        if i not in mapped_indexes and cand['structural_type'] != 'label':
            extractor.unmapped.append({'target': cand['target'], 'reason': 'LLM未映射'})

    # 批次：LLM 草案校验不过 → 默认分批兜底
    batches = llm.get('batches') or []
    if not _batches_valid(batches, set(used_keys)):
        logger.warning('[TemplateAnalyzer] LLM批次草案校验未通过，回退默认分批')
        batches = _default_batches(fields)

    return {
        'schema_version': '2.0',
        'template_name': template_name,
        'doc_stats': {'top_tables': extractor.top_tables,
                      'nested_tables': extractor.nested_tables},
        'fields': fields,
        'batches': batches,
        'unmapped_cells': extractor.unmapped,
    }


# ============================== 对外主入口 ==============================

async def analyze_template(docx_path: str, template_name: str,
                           llm_client: LLMClient,
                           user_hint: str = None, prev_mapping: dict = None) -> tuple[dict, dict]:
    """
    模板分析主入口：结构提取 → LLM 语义映射 → 校验组装

    Args:
        llm_client: LLMClient 实例（settings 主模型，与报告一致）
        user_hint: 用户补充/修正说明（重解析时与上次映射基线一起送 LLM 局部修正）
        prev_mapping: 上次 _semantic_map 的原始输出（重解析基线，首次为 None）

    Returns:
        (落盘 Schema, 本次语义映射原始输出（供下次重解析作基线）)

    Raises:
        InvalidTemplateError: 格式不符（status=invalid）
        Exception: 其他失败（status=failed）
    """
    try:
        doc = Document(docx_path)
    except Exception as e:
        raise InvalidTemplateError(f'文档无法解析(需为有效.docx): {e}') from e

    extractor = _TemplateExtractor().extract(doc)
    # 传给 LLM 的候选清单（含兜底 label，规则8）
    candidates_for_llm = []
    for i, c in enumerate(extractor.candidates):
        cand_llm = {
            'index': i,
            'target': c['target'],
            'structural_type': c['structural_type'],
            'placeholder': c['placeholder'],
            'context': c['context'],
            'options': c['options'],
            'preset_rows': c['preset_rows'],
        }
        if c.get('_anchors'):   # label_group 锚点交给 LLM 逐条覆盖
            cand_llm['anchors'] = c['_anchors']
        if c.get('_blank_label'):   # 选项内嵌填空的语义信息
            cand_llm['blank_label'] = c['_blank_label']
            if c.get('_blank_unit'):
                cand_llm['blank_unit'] = c['_blank_unit']
        candidates_for_llm.append(cand_llm)

    fillable = [c for c in candidates_for_llm if c['structural_type'] != 'label']
    if not fillable:
        raise InvalidTemplateError('未识别到任何可填写字段')

    logger.info('[TemplateAnalyzer] 结构提取完成: 候选%d项(可填%d), 嵌套表%d个',
                len(candidates_for_llm), len(fillable), extractor.nested_tables)
    llm = await _semantic_map(template_name, candidates_for_llm, llm_client,
                              prev_mapping=prev_mapping, user_hint=user_hint)
    return _assemble(template_name, extractor, llm), llm
