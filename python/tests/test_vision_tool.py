"""python/tests/test_vision_tool.py — VisionTool（VLM 图片理解工具）单测。

覆盖：
- 图片直解 / PDF 提图逐张分析 / DOCX 提图
- max_images 截断显式声明（total/truncated/hint，禁止静默 limit）
- 小图过滤显式计数（skipped_small）+ md5 去重
- VLM 未配置前置拦截（VLM_NOT_CONFIGURED）
- 路径越界拒绝 / 不支持格式 / 单张失败不中断批量
- fetched 只读根绝对路径回读 / workspace 未设置

VLM 全部 mock（monkeypatch VLMClient.analyze），不依赖外部服务；
async 用 asyncio.run 驱动（项目惯例，不引入 pytest-asyncio）。
"""
import asyncio
import io
import json
import os
import sys
from pathlib import Path

import pytest

# Ensure python/ is on sys.path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from db.database import Database
from tools.vision import VisionTool


def _run(coro):
    """Synchronously drive a coroutine (project uses plain pytest, no pytest-asyncio)."""
    return asyncio.run(coro)


def _make_png(width: int, height: int) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), (180, 40, 40)).save(buf, format="PNG")
    return buf.getvalue()


def _make_pdf(path: Path, specs: list[tuple[int, int]]):
    """生成含指定尺寸图片的 PDF（每图一页）。"""
    import fitz

    doc = fitz.open()
    for w, h in specs:
        page = doc.new_page(width=595, height=842)
        rect_w = min(w, 400)
        rect_h = min(h, 500)
        page.insert_image(
            fitz.Rect(50, 50, 50 + rect_w, 50 + rect_h),
            stream=_make_png(w, h),
        )
    doc.save(str(path))
    doc.close()


def _make_noisy_png(width: int, height: int) -> bytes:
    """随机噪声 PNG（不可压缩，体积可控地大，用于超大图过滤用例）。"""
    import random

    from PIL import Image

    random.seed(42)  # 确定性：同尺寸同字节，md5 稳定
    buf = io.BytesIO()
    Image.frombytes(
        "RGB", (width, height),
        bytes(random.getrandbits(8) for _ in range(width * height * 3)),
    ).save(buf, format="PNG")
    return buf.getvalue()


def _make_docx(path: Path):
    """生成含一张图片的 DOCX。"""
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph("前文")
    doc.add_picture(io.BytesIO(_make_png(200, 120)), width=Inches(2))
    doc.add_paragraph("后文")
    doc.save(str(path))


@pytest.fixture
def db():
    database = Database(":memory:")
    database.initialize()
    return database


@pytest.fixture
def db_with_vlm(db):
    for key, val in (("vlm_api_url", "https://vlm.example/api"), ("vlm_api_key", "sk-test")):
        db.conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?)", (key, json.dumps(val))
        )
    db.conn.commit()
    return db


@pytest.fixture
def tool(db_with_vlm, tmp_path):
    t = VisionTool(db=db_with_vlm)
    t.set_workspace(str(tmp_path))
    return t


class FakeAnalyze:
    """VLMClient.analyze 替身：记录调用并可脚本化成功/失败序列。

    注意：callable 实例无 descriptor 绑定，client.analyze(path, q) 直接
    映射到 __call__(self, image_path, question)，不经过 self_client。
    """

    def __init__(self, fail_first: int = 0):
        self.calls: list[dict] = []
        self.fail_first = fail_first

    async def __call__(self, image_path, question=None):
        self.calls.append({"path": image_path, "question": question})
        if self.fail_first > 0:
            self.fail_first -= 1
            from errors import VLMRateLimited

            raise VLMRateLimited(detail="429")
        return f"DESC:{Path(image_path).name}"


def _install(monkeypatch, fake: FakeAnalyze):
    monkeypatch.setattr("knowledge.vlm_client.VLMClient.analyze", fake)


class TestImageInput:
    def test_image_returns_description(self, tool, tmp_path, monkeypatch):
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        (tmp_path / "photo.png").write_bytes(_make_png(300, 200))

        result = json.loads(_run(tool.execute(file_path="photo.png")))

        assert result["ok"] is True
        assert result["file"] == "photo.png"
        assert result["description"] == "DESC:photo.png"
        assert len(fake.calls) == 1

    def test_question_forwarded(self, tool, tmp_path, monkeypatch):
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        (tmp_path / "chart.jpg").write_bytes(_make_png(300, 200))

        _run(tool.execute(file_path="chart.jpg", question="趋势是什么？"))

        assert fake.calls[0]["question"] == "趋势是什么？"


class TestDocumentInput:
    def test_pdf_extracts_and_analyzes(self, tool, tmp_path, monkeypatch):
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        _make_pdf(tmp_path / "doc.pdf", [(300, 200), (320, 240)])

        result = json.loads(_run(tool.execute(file_path="doc.pdf")))

        assert result["ok"] is True
        assert result["total_images"] == 2
        assert result["analyzed"] == 2
        assert "truncated" not in result
        assert [e["page"] for e in result["images"]] == [1, 2]
        assert all("description" in e for e in result["images"])
        assert len(fake.calls) == 2

    def test_truncation_is_explicit(self, tool, tmp_path, monkeypatch):
        """图数超上限：显式 truncated/hint，禁止静默 limit（0.1 §16）。"""
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        _make_pdf(tmp_path / "doc.pdf", [(300, 200), (320, 240)])

        result = json.loads(_run(tool.execute(file_path="doc.pdf", max_images=1)))

        assert result["total_images"] == 2
        assert result["analyzed"] == 1
        assert result["truncated"] is True
        assert "hint" in result and "2 张" in result["hint"]
        assert len(fake.calls) == 1

    def test_small_images_skipped_with_count(self, tool, tmp_path, monkeypatch):
        """装饰小图跳过但显式计数，不静默丢弃。"""
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        _make_pdf(tmp_path / "doc.pdf", [(30, 30), (300, 200)])

        result = json.loads(_run(tool.execute(file_path="doc.pdf")))

        assert result["skipped_small"] == 1
        assert result["total_images"] == 1
        assert result["analyzed"] == 1

    def test_duplicate_image_deduped(self, tool, tmp_path, monkeypatch):
        """同一字节图多处引用（页眉 logo）：md5 去重只分析一次。"""
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        png = _make_png(300, 200)

        import fitz

        doc = fitz.open()
        for page_no in range(2):
            page = doc.new_page(width=595, height=842)
            page.insert_image(fitz.Rect(50, 50, 350, 330), stream=png)
        doc.save(str(tmp_path / "dup.pdf"))
        doc.close()

        result = json.loads(_run(tool.execute(file_path="dup.pdf")))

        assert result["total_images"] == 1
        assert len(fake.calls) == 1

    def test_single_failure_does_not_break_batch(self, tool, tmp_path, monkeypatch):
        """批量中单张 VLM 失败：该张记 error 项，其余照常。"""
        fake = FakeAnalyze(fail_first=1)
        _install(monkeypatch, fake)
        _make_pdf(tmp_path / "doc.pdf", [(300, 200), (320, 240)])

        result = json.loads(_run(tool.execute(file_path="doc.pdf")))

        assert result["ok"] is True
        assert result["analyzed"] == 1
        assert "error" in result["images"][0]
        assert "description" in result["images"][1]

    def test_docx_extracts_images(self, tool, tmp_path, monkeypatch):
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        _make_docx(tmp_path / "doc.docx")

        result = json.loads(_run(tool.execute(file_path="doc.docx")))

        assert result["ok"] is True
        assert result["total_images"] == 1
        assert result["analyzed"] == 1


class TestSizeAndCap:
    """review 修复回归：超大图过滤（H1）、max_images 硬上限（M1）、
    office 路径小图/不支持格式过滤（L3/L2）。"""

    def test_oversized_single_image_rejected(self, tool, tmp_path, monkeypatch):
        """单图超限：结构化报错且不触发 VLM，不白传请求体（H1）。"""
        monkeypatch.setattr("tools.vision._MAX_IMAGE_BYTES", 100)
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        (tmp_path / "big.png").write_bytes(_make_png(300, 200))

        result = json.loads(_run(tool.execute(file_path="big.png")))

        assert "图片过大" in result["error"]
        assert "suggestion" in result
        assert fake.calls == []

    def test_oversized_pdf_image_skipped(self, tool, tmp_path, monkeypatch):
        """PDF 内嵌超大图：跳过并显式计数，不占分析预算（H1）。"""
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        monkeypatch.setattr("tools.vision._MAX_IMAGE_BYTES", 10 * 1024)

        import fitz

        doc = fitz.open()
        for stream in (_make_png(64, 64), _make_noisy_png(100, 100)):
            page = doc.new_page(width=595, height=842)
            page.insert_image(fitz.Rect(50, 50, 350, 350), stream=stream)
        doc.save(str(tmp_path / "mixed.pdf"))
        doc.close()

        result = json.loads(_run(tool.execute(file_path="mixed.pdf")))

        assert result["skipped_oversized"] == 1
        assert result["total_images"] == 1
        assert result["analyzed"] == 1

    def test_max_images_hard_cap(self, tool, tmp_path, monkeypatch):
        """max_images 硬上限钳制：无界传参不放大外呼次数（M1）。"""
        fake = FakeAnalyze()
        _install(monkeypatch, fake)
        monkeypatch.setattr("tools.vision._MAX_IMAGES_CAP", 2)
        _make_pdf(tmp_path / "doc.pdf", [(300, 200), (320, 240), (340, 260)])

        result = json.loads(_run(tool.execute(file_path="doc.pdf", max_images=999)))

        assert result["total_images"] == 3
        assert result["analyzed"] == 2
        assert result["truncated"] is True
        assert len(fake.calls) == 2

    def test_docx_small_image_skipped(self, tool, tmp_path, monkeypatch):
        """DOCX 小图标跳过并显式计数，与 PDF 路径一致（L3）。"""
        fake = FakeAnalyze()
        _install(monkeypatch, fake)

        from docx import Document
        from docx.shared import Inches

        doc = Document()
        doc.add_picture(io.BytesIO(_make_png(30, 30)), width=Inches(0.3))
        doc.add_picture(io.BytesIO(_make_png(200, 120)), width=Inches(2))
        doc.save(str(tmp_path / "mixed.docx"))

        result = json.loads(_run(tool.execute(file_path="mixed.docx")))

        assert result["skipped_small"] == 1
        assert result["total_images"] == 1
        assert result["analyzed"] == 1

    def test_docx_tiff_image_skipped(self, tool, tmp_path, monkeypatch):
        """DOCX 内嵌 tif（office 白名单放行、VLM MIME 不支持）：显式计数跳过。"""
        fake = FakeAnalyze()
        _install(monkeypatch, fake)

        from PIL import Image
        from docx import Document
        from docx.shared import Inches

        tif_buf = io.BytesIO()
        Image.new("RGB", (300, 200), (40, 40, 180)).save(tif_buf, format="TIFF")
        doc = Document()
        doc.add_picture(io.BytesIO(tif_buf.getvalue()), width=Inches(2))
        doc.save(str(tmp_path / "pic.docx"))

        result = json.loads(_run(tool.execute(file_path="pic.docx")))

        assert result["skipped_unsupported"] == 1
        assert result["total_images"] == 0
        assert result["analyzed"] == 0


class TestGuards:
    def test_vlm_not_configured(self, tmp_path, monkeypatch):
        """VLM 未配置：前置拦截返回结构化错误，不触发提图。"""
        db = Database(":memory:")
        db.initialize()
        tool = VisionTool(db=db)
        tool.set_workspace(str(tmp_path))
        _make_pdf(tmp_path / "doc.pdf", [(300, 200)])

        result = json.loads(_run(tool.execute(file_path="doc.pdf")))

        assert result["code"] == "VLM_NOT_CONFIGURED"
        assert "设置" in result["suggestion"]

    def test_path_traversal_blocked(self, tool, monkeypatch):
        fake = FakeAnalyze()
        _install(monkeypatch, fake)

        result = json.loads(_run(tool.execute(file_path="../outside.png")))

        assert "error" in result
        assert fake.calls == []

    def test_missing_file(self, tool):
        result = json.loads(_run(tool.execute(file_path="nope.png")))
        assert "error" in result and "不存在" in result["error"]

    def test_unsupported_format(self, tool, tmp_path):
        (tmp_path / "note.txt").write_text("hello", encoding="utf-8")

        result = json.loads(_run(tool.execute(file_path="note.txt")))

        assert "error" in result
        assert ".pdf" in result["supported"]

    def test_workspace_not_set(self, db_with_vlm, monkeypatch):
        tool = VisionTool(db=db_with_vlm)
        fake = FakeAnalyze()
        _install(monkeypatch, fake)

        result = json.loads(_run(tool.execute(file_path="x.png")))

        assert "error" in result and "工作目录" in result["error"]

    def test_fetched_readonly_root_absolute_path(self, db_with_vlm, tmp_path, monkeypatch):
        """只读根内绝对路径可读（对齐 file_ops fetched 语义）。"""
        fetched = tmp_path / "fetched_root"
        fetched.mkdir()
        (fetched / "shot.png").write_bytes(_make_png(300, 200))

        tool = VisionTool(db=db_with_vlm)
        tool.set_workspace(str(tmp_path / "ws"))
        tool.add_readonly_root(str(fetched))
        fake = FakeAnalyze()
        _install(monkeypatch, fake)

        result = json.loads(_run(tool.execute(file_path=str(fetched / "shot.png"))))

        assert result["ok"] is True
        assert result["description"] == "DESC:shot.png"

    def test_outside_readonly_root_denied(self, db_with_vlm, tmp_path, monkeypatch):
        """未注册的绝对路径拒绝（沙箱不因只读根放开）。"""
        outside = tmp_path / "outside"
        outside.mkdir()
        (outside / "x.png").write_bytes(_make_png(300, 200))

        tool = VisionTool(db=db_with_vlm)
        tool.set_workspace(str(tmp_path / "ws"))
        fake = FakeAnalyze()
        _install(monkeypatch, fake)

        result = json.loads(_run(tool.execute(file_path=str(outside / "x.png"))))

        assert "error" in result
