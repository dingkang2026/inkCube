# -*- coding: utf-8 -*-
"""form_fill.fill_service 单元测试。

覆盖：批次拓扑排序（依赖序/环兜底）、_normalize_value 严校验与幂等（人工修订
通道共用）、_merge_queries 去重限量、_digest 摘要、联合检索（项目库 + 基础库
去重 + 失败降级）、fill_template 端到端（FakeLLMClient + FakeSearchTool，
断言产物 docx 内容 / 报告 / 值集 / 进度回调）。
"""
import asyncio
import io
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document

from form_fill.fill_service import (
    _digest,
    _merge_queries,
    _normalize_value,
    _order_batches,
    _retrieve,
    fill_template,
)
from form_fill.prompts import TEMPLATE_FILL_PROMPT, build_fill_user_prompt


# ============================== 桩 ==============================

class FakeLLMClient:
    """LLMClient 桩：按调用序返回固定值集 JSON；元素为 Exception 实例时抛出"""

    def __init__(self, contents):
        self._contents = list(contents)
        self.calls = []

    async def chat_full(self, messages):
        self.calls.append({"messages": messages})
        item = self._contents.pop(0)
        if isinstance(item, Exception):
            raise item
        return {"content": item, "finish_reason": "stop",
                "usage": {"prompt_tokens": 1, "completion_tokens": 1,
                          "total_tokens": 2}}


class FakeSearchTool:
    """KnowledgeSearchTool 桩：按 (query, project_id) 返回固定结果并记录调用"""

    def __init__(self, results_by_query=None):
        self.results_by_query = results_by_query or {}
        self.calls = []   # [(query, project_id)]

    async def search(self, query, project_id, top_k=5, **kwargs):
        self.calls.append((query, project_id))
        return list(self.results_by_query.get(query, []))[:top_k]


# ============================== 批次拓扑排序 ==============================

def test_order_batches_respects_dependencies():
    batches = [
        {"seq": 2, "field_keys": ["b"], "depends_on": [1]},
        {"seq": 1, "field_keys": ["a"], "depends_on": []},
        {"seq": 3, "field_keys": ["c"], "depends_on": [1, 2]},
    ]
    ordered = _order_batches(batches)
    assert [b["seq"] for b in ordered] == [1, 2, 3]


def test_order_batches_cycle_fallback_by_seq():
    """依赖成环：按 seq 兜底全收（不丢批次）"""
    batches = [
        {"seq": 2, "field_keys": ["b"], "depends_on": [1]},
        {"seq": 1, "field_keys": ["a"], "depends_on": [2]},
    ]
    ordered = _order_batches(batches)
    assert len(ordered) == 2
    assert [b["seq"] for b in ordered] == [1, 2]


def test_order_batches_unknown_dep_treated_as_done():
    """depends_on 引用不存在的 seq：视作已满足（不阻塞）"""
    batches = [{"seq": 1, "field_keys": ["a"], "depends_on": [99]}]
    assert [b["seq"] for b in _order_batches(batches)] == [1]


# ============================== 值归一化（含幂等） ==============================

def _choice_field(**over):
    field = {"field_key": "industry", "label": "所属行业", "fill_type": "choice",
             "target": "top0/r0c1", "options": ["人工智能", "高端装备"],
             "multi_select": False}
    field.update(over)
    return field


def test_normalize_value_text_and_blank():
    text_field = {"field_key": "name", "fill_type": "text", "target": "top0/r0c1"}
    assert _normalize_value(text_field, "  星尘科技 ") == "星尘科技"
    assert _normalize_value(text_field, "   ") is None
    assert _normalize_value(text_field, None) is None

    blank_field = {"field_key": "amount", "fill_type": "text", "target": "top0/r1c1#blank1",
                   "unit": "万元"}
    assert _normalize_value(blank_field, "1000万元") == "1000"
    assert _normalize_value(blank_field, "约1000.5") == "1000.5"
    # 无数字且非空 → 保留原文（仅空串归 None）；空串/None → None
    assert _normalize_value(blank_field, "—") == "—"
    assert _normalize_value(blank_field, "") is None


def test_normalize_value_choice_whitelist_and_other():
    field = _choice_field()
    # 白名单内
    v = _normalize_value(field, {"selected": ["人工智能"], "other": None})
    assert v == {"selected": ["人工智能"], "other": None}
    # 白名单外 → 落 other
    v = _normalize_value(field, {"selected": ["量子计算"], "other": None})
    assert v == {"selected": [], "other": "量子计算"}
    # 单选收敛（multi_select=False）
    v = _normalize_value(field, {"selected": ["人工智能", "高端装备"], "other": None})
    assert v == {"selected": ["人工智能"], "other": None}
    # 空值集 → None（记 no_data）
    assert _normalize_value(field, {"selected": [], "other": None}) is None
    # 裸字符串入参（容错）
    v = _normalize_value(field, "高端装备")
    assert v == {"selected": ["高端装备"], "other": None}


def test_normalize_value_label_group_and_table_rows():
    lg_field = {"field_key": "deal", "fill_type": "label_group", "target": "top0/r0c1"}
    assert _normalize_value(lg_field, {"a": "1"}) == {"a": "1"}
    assert _normalize_value(lg_field, {}) is None
    assert _normalize_value(lg_field, "str") is None

    tr_field = {"field_key": "rows", "fill_type": "table_rows", "target": "top0/r0c1"}
    assert _normalize_value(tr_field, [["a", 1], [None, "b"]]) == [["a", "1"], ["", "b"]]
    assert _normalize_value(tr_field, "not-a-list") is None
    assert _normalize_value(tr_field, []) is None


def test_normalize_value_idempotent():
    """对自身输出再归一一次结果不变（人工修订通道共用，设计 §2.2）"""
    fields = [
        {"field_key": "amount", "fill_type": "text", "target": "top0/r1c1#blank1",
         "unit": "万元"},
        _choice_field(),
        _choice_field(multi_select=True),
        {"field_key": "deal", "fill_type": "label_group", "target": "top0/r0c1"},
        {"field_key": "rows", "fill_type": "table_rows", "target": "top0/r0c1"},
        {"field_key": "name", "fill_type": "text", "target": "top0/r0c1"},
    ]
    raw_values = [
        "1000万元", {"selected": ["人工智能"], "other": None},
        {"selected": ["人工智能", "高端装备", "量子计算"], "other": "x"},
        {"a": "1"}, [["a", 1], [None, "b"]], "  星尘科技 ",
    ]
    for field, raw in zip(fields, raw_values):
        once = _normalize_value(field, raw)
        twice = _normalize_value(field, once)
        assert twice == once, f"non-idempotent for {field['field_key']}: {once!r} -> {twice!r}"


# ============================== 检索词合并 ==============================

def test_merge_queries_dedup_and_limit():
    """字段级检索词去重收集（merged_kb_queries 层已删除，P1）"""
    fields = [
        {"kb_query": ["所属行业", "主营业务"],
         "sub_fields": [{"kb_query": ["主营业务", "出资基金"]}]},
        {"kb_query": []},
    ]
    queries = _merge_queries(fields)
    assert queries == ["所属行业", "主营业务", "出资基金"]


def test_merge_queries_round_robin_keeps_every_field():
    """轮转交错：限量截断后每个字段至少保留首条检索词（不饿死靠后字段）"""
    fields = [
        {"kb_query": [f"{p}{i}" for i in range(6)]}
        for p in "abcde"
    ]
    queries = _merge_queries(fields)
    assert len(queries) == 8
    # 每字段首条都进入（轮转第 1 轮 5 条 + 第 2 轮 3 条）
    for first in ("a0", "b0", "c0", "d0", "e0"):
        assert first in queries
    # 交错序：第 2 轮按字段序续接（a1 b1 c1）
    assert queries[5:] == ["a1", "b1", "c1"]


def test_merge_queries_batch_wider_than_limit_keeps_every_first():
    """字段数超 QUERIES_PER_BATCH（analyzer 允许至 15）：保底下限随字段数
    扩展，每个字段首条无条件保留（review 发现：硬限量会饿死第 9+ 字段）"""
    fields = [
        {"kb_query": [f"{p}{i}" for i in range(2)]}
        for p in "abcdefghij"   # 10 字段 × 2 词
    ]
    queries = _merge_queries(fields)
    assert len(queries) == 10
    assert queries == [f"{p}0" for p in "abcdefghij"]   # 第 1 轮全员保底


# ============================== 摘要 ==============================

def test_digest_previews():
    fields = {"a": {"label": "名称"}, "rows": {"label": "表格"}}
    digest = _digest({"a": "星尘科技", "rows": [["x"], ["y"]]}, fields)
    lines = digest.split("\n")
    assert lines[0] == "- a(名称): 星尘科技"
    assert lines[1] == "- rows(表格): 2行"


# ============================== 联合检索 ==============================

def test_retrieve_joins_projects_and_dedups():
    """项目库 + 基础库逐组合检索；跨库重复 chunk 按 id 去重；按分排序限量"""
    tool = FakeSearchTool()

    async def search(query, project_id, top_k=5, **kw):
        tool.calls.append((query, project_id))
        if project_id == "proj_1":
            return [{"id": "c1", "content": "项目库内容", "_score": 1.0}]
        return [{"id": "c1", "content": "基础库同块", "_score": 5.0},
                {"id": "c2", "content": "基础库新增", "_score": 2.0}]

    tool.search = search
    context = asyncio.run(_retrieve(["主营业务"], tool, "proj_1", ["lib_1"]))
    # 两个 project_id 均被检索
    assert ("主营业务", "proj_1") in tool.calls
    assert ("主营业务", "lib_1") in tool.calls
    # 同 id 去重（首见保留：项目库内容），另含基础库新增块
    assert "项目库内容" in context
    assert "基础库新增" in context
    assert "基础库同块" not in context
    assert context.startswith("【检索词：主营业务】")


def test_retrieve_failure_degrades_to_empty():
    """检索抛异常 → 该词降级为空，不影响其他词"""
    tool = FakeSearchTool()

    async def search(query, project_id, top_k=5, **kw):
        raise RuntimeError("embedding model not ready")

    tool.search = search
    context = asyncio.run(_retrieve(["q1"], tool, "proj_1", []))
    assert context == ""


def test_retrieve_source_annotation():
    """带 metadata.filename 的块追加来源行；无 metadata 块不加（主体归属甄别依据）"""
    tool = FakeSearchTool()

    async def search(query, project_id, top_k=5, **kw):
        return [
            {"id": "c1", "content": "营业执照内容", "_score": 2.0,
             "metadata": {"filename": "湖南宁庆营业执照.jpg"}},
            {"id": "c2", "content": "无来源块内容", "_score": 1.0},   # 无 metadata
        ]

    tool.search = search
    context = asyncio.run(_retrieve(["营业执照"], tool, "proj_1", []))
    assert "营业执照内容\n—— 来源：湖南宁庆营业执照.jpg" in context
    assert "无来源块内容" in context
    assert "无来源块内容\n—— 来源" not in context


def test_retrieve_dedups_chunks_across_queries():
    """跨检索词去重（P1）：同一 chunk 只进首个命中的词的 block，不重复占预算"""
    tool = FakeSearchTool()

    async def search(query, project_id, top_k=5, **kw):
        if query == "企业全称":
            return [{"id": "c1", "content": "全称资料", "_score": 1.0}]
        return [{"id": "c1", "content": "全称资料", "_score": 2.0},   # 与前一词撞块
                {"id": "c2", "content": "行业资料", "_score": 1.0}]

    tool.search = search
    context = asyncio.run(_retrieve(["企业全称", "所属行业"], tool, "proj_1", []))
    assert context.count("全称资料") == 1
    assert "行业资料" in context


# ============================== Prompt 组装（项目主体锚定） ==============================

def test_build_fill_user_prompt_subject_anchor():
    """主体声明块：名称/信用代码齐全，位于检索资料之后、基础信息之前，含甄别指令"""
    prompt = build_fill_user_prompt(
        "投资推荐项目情况表",
        [{"field_key": "office_address", "label": "办公地址", "fill_type": "text"}],
        "【检索词：办公地址】\n某公司地址资料",
        {"company_name": "宁庆空天", "credit_code": "91320000TEST", "today": "2026-10-15"},
        "", kb_chars_max=24000)
    subject_part = prompt.split("## 项目主体")[1].split("##")[0]
    assert "- 名称：宁庆空天" in subject_part
    assert "- 统一社会信用代码：91320000TEST" in subject_part
    assert "不得作为字段填写依据" in subject_part
    assert prompt.index("## 项目主体") < prompt.index("## 项目基础信息")


def test_build_fill_user_prompt_subject_degrades_when_unregistered():
    """基础信息未登记：主体段降级（占位说明 + 通用甄别指令照给）"""
    prompt = build_fill_user_prompt("t", [], "", {}, "", kb_chars_max=24000)
    assert "## 项目主体" in prompt
    assert "未登记名称与信用代码" in prompt
    assert "不得作为字段填写依据" in prompt


def test_fill_system_prompt_has_subject_discipline():
    """系统 prompt 资料纪律含主体纪律（与 user_prompt 甄别要求呼应）"""
    assert "项目主体" in TEMPLATE_FILL_PROMPT
    assert "不得作为字段依据" in TEMPLATE_FILL_PROMPT


# ============================== fill_template 端到端 ==============================


def _build_form_doc():
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("项目名称")
    table.cell(1, 0).paragraphs[0].add_run("所属行业")
    table.cell(1, 1).paragraphs[0].add_run("□人工智能 □高端装备")
    table.cell(2, 0).paragraphs[0].add_run("申请日期")
    return doc


_SCHEMA = {
    "schema_version": "2.0",
    "template_name": "测试模板",
    "fields": [
        {"field_key": "project_name", "label": "项目名称", "fill_type": "text",
         "fill_mode": "kb_fill", "target": "top0/r0c1", "placeholder": None,
         "context": {}, "kb_query": ["项目名称"], "confidence": "high"},
        {"field_key": "industry", "label": "所属行业", "fill_type": "choice",
         "fill_mode": "kb_fill", "target": "top0/r1c1", "placeholder": None,
         "context": {}, "kb_query": ["所属行业"], "confidence": "high",
         "options": ["人工智能", "高端装备"], "multi_select": False},
        {"field_key": "apply_date", "label": "申请日期", "fill_type": "text",
         "fill_mode": "rule_fill", "target": "top0/r2c1", "placeholder": None,
         "context": {}, "kb_query": [], "confidence": "high",
         "rule_hint": "today"},
    ],
    "batches": [
        {"seq": 1, "name": "头部信息",
         "field_keys": ["project_name", "industry"], "depends_on": []},
        {"seq": 2, "name": "落款", "field_keys": ["apply_date"], "depends_on": []},
    ],
    "unmapped_cells": [],
}


def test_fill_template_end_to_end():
    llm_values = {"project_name": "星尘科技", "industry": {"selected": ["人工智能"], "other": None}}
    llm = FakeLLMClient([json.dumps(llm_values, ensure_ascii=False)])
    tool = FakeSearchTool(results_by_query={
        "项目名称": [{"id": "c1", "content": "企业全称：星尘科技", "_score": 1.0}],
    })
    progress = []

    with tempfile.TemporaryDirectory() as td:
        doc_path = os.path.join(td, "tpl.docx")
        _build_form_doc().save(doc_path)
        doc_bytes, report, filled = asyncio.run(fill_template(
            _SCHEMA, doc_path,
            {"company_name": "星尘科技", "credit_code": "91320000TEST",
             "apply_date": "2026-10-14", "today": "2026-10-14"},
            tool, "proj_1", ["lib_1"], llm, {"api_key": "k"},
            progress_cb=lambda done, total: progress.append((done, total)),
        ))

    # 产物：text 直填 + 勾选替换 + rule_fill 直填
    out = Document(io.BytesIO(doc_bytes))
    assert out.tables[0].cell(0, 1).text == "星尘科技"
    assert "☑人工智能" in out.tables[0].cell(1, 1).text
    assert "□高端装备" in out.tables[0].cell(1, 1).text
    assert out.tables[0].cell(2, 1).text == "2026-10-14"

    # 报告：三字段均 filled；值集含三键
    by_key = {r["field_key"]: r for r in report}
    assert by_key["project_name"]["status"] == "filled"
    assert by_key["industry"]["status"] == "filled"
    assert by_key["apply_date"]["status"] == "filled"
    assert set(filled) == {"project_name", "industry", "apply_date"}

    # 检索只发生在 kb_fill 批次（批 1），且项目库 + 基础库都被检索
    assert any(pid == "proj_1" for _, pid in tool.calls)
    assert any(pid == "lib_1" for _, pid in tool.calls)

    # LLM 只调用 1 次（批 2 全 rule_fill，无 LLM 调用）
    assert len(llm.calls) == 1
    assert llm.calls[0]["messages"][0]["role"] == "system"

    # 主体锚定端到端：user_prompt 含主体段（名称来自 project_context）
    user_prompt = llm.calls[0]["messages"][1]["content"]
    assert "## 项目主体" in user_prompt
    assert "- 名称：星尘科技" in user_prompt

    # 进度回调：逐批单调推进
    assert progress == [(1, 2), (2, 2)]


def test_fill_template_no_data_keeps_blank():
    """查无资料 → 留空不写占位文字，报告记 no_data（宁缺毋假）"""
    llm = FakeLLMClient([json.dumps({"project_name": None, "industry": None},
                                    ensure_ascii=False)])
    tool = FakeSearchTool()   # 无任何资料
    with tempfile.TemporaryDirectory() as td:
        doc_path = os.path.join(td, "tpl.docx")
        _build_form_doc().save(doc_path)
        schema = dict(_SCHEMA)
        doc_bytes, report, filled = asyncio.run(fill_template(
            schema, doc_path, {"apply_date": "2026-10-14"},
            tool, "proj_1", [], llm, {"api_key": "k"},
        ))
    out = Document(io.BytesIO(doc_bytes))
    assert out.tables[0].cell(0, 1).text == ""           # 留空
    assert "☑" not in out.tables[0].cell(1, 1).text      # 不勾选
    by_key = {r["field_key"]: r for r in report}
    assert by_key["project_name"]["status"] == "no_data"
    assert by_key["industry"]["status"] == "no_data"
    assert filled == {"apply_date": "2026-10-14"}


def test_fill_template_batch_llm_failure_isolated():
    """单批 LLM 失败 → 该批字段记 failed，不中断后续批次"""
    llm = FakeLLMClient([RuntimeError("boom"), "{}"])   # 批 1 异常、批 2 空值集
    tool = FakeSearchTool()
    with tempfile.TemporaryDirectory() as td:
        doc_path = os.path.join(td, "tpl.docx")
        _build_form_doc().save(doc_path)
        schema = {
            "schema_version": "2.0", "template_name": "t",
            "fields": [
                {"field_key": "project_name", "label": "项目名称", "fill_type": "text",
                 "fill_mode": "kb_fill", "target": "top0/r0c1", "placeholder": None,
                 "context": {}, "kb_query": ["q"], "confidence": "high"},
                {"field_key": "industry", "label": "所属行业", "fill_type": "text",
                 "fill_mode": "kb_fill", "target": "top0/r1c1", "placeholder": None,
                 "context": {}, "kb_query": ["q2"], "confidence": "high"},
            ],
            "batches": [
                {"seq": 1, "name": "b1", "field_keys": ["project_name"], "depends_on": []},
                {"seq": 2, "name": "b2", "field_keys": ["industry"], "depends_on": []},
            ],
            "unmapped_cells": [],
        }
        doc_bytes, report, filled = asyncio.run(fill_template(
            schema, doc_path, {}, tool, "proj_1", [], llm,
            {"api_key": "k"},
        ))
    by_key = {r["field_key"]: r for r in report}
    assert by_key["project_name"]["status"] == "failed"   # 批 1 LLM 异常
    assert by_key["industry"]["status"] == "no_data"      # 批 2 正常执行（LLM 返回空值集）
    assert "industry" not in filled
