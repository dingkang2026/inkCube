# -*- coding: utf-8 -*-
"""form_fill.docx_filler 单元测试。

覆盖：XML 路径解析（top/nested/#blank）、清格写入（保留字体）、勾选替换 3 种
跨 run 分布、label_group 锚点插值、#blank 定位（标签优先/序号兜底/其他排除）、
table_rows 插行、apply_field 分发。
异步用例统一用 asyncio.run 驱动（本模块无异步，约定见 test_card_builder.py）。
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from docx.oxml.ns import qn

from form_fill.docx_filler import (
    apply_field,
    check_options,
    fill_after_label,
    fill_blank,
    fill_table,
    resolve_node,
    write_text,
)


def _doc_with_grid(rows: int, cols: int):
    doc = Document()
    table = doc.add_table(rows=rows, cols=cols)
    return doc, table


def _set_cell(table, r, c, text):
    """单 run 写入单元格（默认段落已有，add_run 追加）"""
    cell = table.cell(r, c)
    cell.paragraphs[0].add_run(text)
    return cell


def _cell_text(table, r, c):
    return table.cell(r, c).text


# ============================== 路径解析 ==============================

def test_resolve_node_top_cell():
    doc, table = _doc_with_grid(2, 2)
    _set_cell(table, 1, 1, "值格")
    node, blank_n = resolve_node(doc, "top0/r1c1")
    assert blank_n is None
    texts = ["".join(t.text or "" for t in p.findall(".//" + qn("w:t")))
             for p in node.findall(qn("w:p"))]
    assert "值格" in texts


def test_resolve_node_blank_suffix():
    doc, _ = _doc_with_grid(1, 1)
    node, blank_n = resolve_node(doc, "top0/r0c0#blank2")
    assert node is not None
    assert blank_n == 2


def test_resolve_node_nested_table():
    doc, table = _doc_with_grid(1, 1)
    inner = table.cell(0, 0).add_table(rows=2, cols=2)
    inner.cell(1, 1).paragraphs[0].add_run("内层")
    node, blank_n = resolve_node(doc, "top0/r0c0/nested0")
    assert blank_n is None
    trs = node.findall(qn("w:tr"))
    assert len(trs) == 2
    tcs = trs[1].findall(qn("w:tc"))
    assert "内层" in "".join(t.text or "" for t in tcs[1].findall(".//" + qn("w:t")))


# ============================== 清格写入 ==============================

def test_write_text_clears_and_multiline():
    doc, table = _doc_with_grid(1, 1)
    cell = _set_cell(table, 0, 0, "旧内容一")
    cell.add_paragraph("旧内容二")
    tc = table.cell(0, 0)._tc
    write_text(tc, "第一段\n第二段")
    texts = [p.text for p in table.cell(0, 0).paragraphs]
    assert texts == ["第一段", "第二段"]


def test_write_text_keeps_first_run_font():
    doc, table = _doc_with_grid(1, 1)
    run = table.cell(0, 0).paragraphs[0].add_run("占位")
    run.font.name = "宋体"
    tc = table.cell(0, 0)._tc
    write_text(tc, "新值")
    new_run = table.cell(0, 0).paragraphs[0].runs[0]
    assert new_run.text == "新值"
    assert new_run.font.name == "宋体"


# ============================== 勾选替换（3 种跨 run 分布） ==============================

def test_check_options_same_run():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "□人工智能 □高端装备")
    tc = table.cell(0, 0)._tc
    missed = check_options(tc, ["人工智能"])
    assert missed == []
    text = _cell_text(table, 0, 0)
    assert "☑人工智能" in text
    assert "□高端装备" in text


def test_check_options_symbol_alone_run():
    """情况2：□ 在前 run 末尾、选项在下 run 开头"""
    doc, table = _doc_with_grid(1, 1)
    cell = table.cell(0, 0)
    p = cell.paragraphs[0]
    p.add_run("□ ")
    p.add_run("人工智能")
    p.add_run(" ")
    p.add_run("□高端装备")
    tc = cell._tc
    missed = check_options(tc, ["人工智能"])
    assert missed == []
    text = _cell_text(table, 0, 0)
    assert "☑" in text and "人工智能" in text
    assert "□高端装备" in text


def test_check_options_prefix_same_run_cross_run():
    """情况3：□ 与选项前缀同 run、选项其余跨 run"""
    doc, table = _doc_with_grid(1, 1)
    cell = table.cell(0, 0)
    p = cell.paragraphs[0]
    p.add_run("□人工")
    p.add_run("智能 □高端装备")
    tc = cell._tc
    missed = check_options(tc, ["人工智能"])
    assert missed == []
    text = _cell_text(table, 0, 0)
    assert "☑人工" in text
    assert "□高端装备" in text


def test_check_options_quote_normalize():
    """选项带中文引号时引号样式归一命中"""
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "□\u201c人工智能\u201d产业")
    tc = table.cell(0, 0)._tc
    missed = check_options(tc, ['"人工智能"产业'])
    assert missed == []
    assert "☑" in _cell_text(table, 0, 0)


def test_check_options_multi_and_missed():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "□人工智能 □高端装备")
    tc = table.cell(0, 0)._tc
    missed = check_options(tc, ["人工智能", "高端装备", "不存在项"])
    assert missed == ["不存在项"]
    text = _cell_text(table, 0, 0)
    assert "☑人工智能" in text and "☑高端装备" in text


# ============================== label_group 锚点插值 ==============================

def test_fill_after_label_replaces_tail():
    doc, table = _doc_with_grid(1, 1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].add_run("出资基金：星尘一期（提示文案）")
    cell.add_paragraph("投资方式：")
    tc = cell._tc
    assert fill_after_label(tc, "出资基金", "XX创投基金") is True
    texts = [p.text for p in cell.paragraphs]
    assert texts[0] == "出资基金：XX创投基金"
    assert texts[1] == "投资方式："


def test_fill_after_label_halfwidth_colon():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "investor: ")
    tc = table.cell(0, 0)._tc
    assert fill_after_label(tc, "investor", "XYZ") is True
    # 冒号后整段（含原空格）被替换为 value
    assert _cell_text(table, 0, 0) == "investor:XYZ"


def test_fill_after_label_anchor_with_hint_text():
    """anchor 可能是"标签：提示"全串——按冒号前主体匹配"""
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "公司概况：请填写公司整体介绍")
    tc = table.cell(0, 0)._tc
    assert fill_after_label(tc, "公司概况：请填写公司整体介绍", "简介内容") is True
    assert _cell_text(table, 0, 0) == "公司概况：简介内容"


def test_fill_after_label_miss():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "无锚点内容")
    tc = table.cell(0, 0)._tc
    assert fill_after_label(tc, "不存在", "值") is False


# ============================== #blank 定位 ==============================

def test_fill_blank_label_priority_and_unit_kept():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "融资金额：____万元；投前估值：____亿元")
    tc = table.cell(0, 0)._tc
    assert fill_blank(tc, 1, "1000", label="融资金额(万元)") is True
    text = _cell_text(table, 0, 0)
    assert "融资金额：1000万元" in text
    assert "投前估值：____亿元" in text


def test_fill_blank_index_fallback_skips_other():
    """无标签匹配时按序号取空白位，且"其他"位不参与编号"""
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "注册资本：____；其他：____")
    tc = table.cell(0, 0)._tc
    assert fill_blank(tc, 1, "500万", label=None) is True
    text = _cell_text(table, 0, 0)
    assert "注册资本：500万" in text
    assert "其他：____" in text


def test_fill_blank_no_match():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "无空白位")
    tc = table.cell(0, 0)._tc
    assert fill_blank(tc, 1, "值", label=None) is False


# ============================== table_rows 插行 ==============================

def test_fill_table_inserts_rows_when_short():
    doc, table = _doc_with_grid(2, 2)
    _set_cell(table, 0, 0, "姓名")
    _set_cell(table, 0, 1, "年龄")   # r1 为空数据行模板
    fill_table(table._tbl, [["张三", "30"], ["李四", "25"]], header_rows=1)
    assert _cell_text(table, 1, 0) == "张三"
    assert _cell_text(table, 1, 1) == "30"
    assert _cell_text(table, 2, 0) == "李四"
    assert _cell_text(table, 2, 1) == "25"


def test_fill_table_skips_empty_values():
    doc, table = _doc_with_grid(2, 2)
    _set_cell(table, 0, 0, "姓名")
    _set_cell(table, 0, 1, "年龄")
    _set_cell(table, 1, 1, "旧值")   # 值列为 None/空 → 该格已有内容不覆盖
    fill_table(table._tbl, [["张三", None]], header_rows=1)
    assert _cell_text(table, 1, 0) == "张三"
    assert _cell_text(table, 1, 1) == "旧值"


# ============================== apply_field 分发 ==============================

def test_apply_field_text_and_blank_dispatch():
    doc, table = _doc_with_grid(1, 2)
    _set_cell(table, 0, 0, "项目名称")
    field = {"field_key": "project_name", "label": "项目名称",
             "fill_type": "text", "target": "top0/r0c1"}
    result = apply_field(doc, field, "星尘科技")
    assert result["ok"] is True
    assert _cell_text(table, 0, 1) == "星尘科技"


def test_apply_field_choice_with_other():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "□人工智能 □高端装备 其他：____")
    field = {"field_key": "industry", "label": "所属行业",
             "fill_type": "choice", "target": "top0/r0c0",
             "options": ["人工智能", "高端装备", "其他"]}
    result = apply_field(doc, field, {"selected": ["高端装备"], "other": "量子计算"})
    assert result["missed"] == []
    text = _cell_text(table, 0, 0)
    assert "☑高端装备" in text
    assert "□人工智能" in text
    assert "其他：量子计算" in text


def test_apply_field_label_group_uses_sub_anchors():
    doc, table = _doc_with_grid(1, 1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].add_run("出资基金：")
    cell.add_paragraph("投资方式：")
    field = {
        "field_key": "deal_info", "label": "交易信息",
        "fill_type": "label_group", "target": "top0/r0c0",
        "sub_fields": [
            {"sub_key": "fund_name", "anchor": "出资基金", "kb_query": []},
            {"sub_key": "invest_way", "anchor": "投资方式", "kb_query": []},
        ],
    }
    result = apply_field(doc, field, {"fund_name": "XX基金", "invest_way": "增资"})
    assert result["ok"] is True and result["missed"] == []
    texts = [p.text for p in cell.paragraphs]
    assert texts[0] == "出资基金：XX基金"
    assert texts[1] == "投资方式：增资"


def test_apply_field_label_group_miss_records():
    doc, table = _doc_with_grid(1, 1)
    _set_cell(table, 0, 0, "出资基金：")
    field = {
        "field_key": "deal_info", "label": "交易信息",
        "fill_type": "label_group", "target": "top0/r0c0",
        "sub_fields": [{"sub_key": "fund_name", "anchor": "出资基金", "kb_query": []}],
    }
    result = apply_field(doc, field, {"fund_name": "XX", "ghost": "yy"})
    assert result["missed"] == ["ghost"]


def test_apply_field_table_rows_on_top_table():
    doc, table = _doc_with_grid(2, 2)
    _set_cell(table, 0, 0, "姓名")
    _set_cell(table, 0, 1, "年龄")
    field = {"field_key": "members", "label": "核心成员",
             "fill_type": "table_rows", "target": "top0", "header_rows": 1}
    result = apply_field(doc, field, [["张三", "30"]])
    assert result["ok"] is True
    assert _cell_text(table, 1, 0) == "张三"


def test_apply_field_doc_section_multiline():
    doc, table = _doc_with_grid(1, 2)
    _set_cell(table, 0, 0, "公司简介")
    field = {"field_key": "profile", "label": "公司简介",
             "fill_type": "doc_section", "target": "top0/r0c1"}
    apply_field(doc, field, "一、主营情况\n\n二、竞争优势")
    texts = [p.text for p in table.cell(0, 1).paragraphs]
    assert texts == ["一、主营情况", "", "二、竞争优势"]


def test_apply_field_none_value_noop_for_text():
    doc, table = _doc_with_grid(1, 2)
    _set_cell(table, 0, 0, "项目名称")
    field = {"field_key": "project_name", "label": "项目名称",
             "fill_type": "text", "target": "top0/r0c1"}
    result = apply_field(doc, field, None)
    assert result["ok"] is True
    assert _cell_text(table, 0, 1) == ""
