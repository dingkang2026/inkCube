"""python/tests/test_office_media.py

DOCX/PPTX 原生图片无损提取(2.2 §4.4 图片路,Office 图片无损改造):
- 字节级无损:提取产物 md5 == media 原始字节 md5(核心断言)
- PPTX:页码 = slide 序号、group 内图、占位符图(PlaceholderPicture)、同 slide 文本上下文
- DOCX:inline + 浮动(anchor)drawing、文档序上下文、页码 None、表格内图
- 按内容 md5 去重;不可读格式跳过不残留;成功后孤儿对账清理;损坏容器异常上抛
"""
import base64
import hashlib
import io
from pathlib import Path

import pytest

from knowledge.office_media import (
    _CONTEXT_BLOCKS,
    _save_and_measure,
    extract_office_images,
)

# 颜色不同 → 字节不同 → 不互相去重
_PNG_RED = None
_PNG_BLUE = None
_PNG_GREEN = None


def _png_of(color):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (320, 240), color).save(buf, format="PNG")
    return buf.getvalue()


def setup_module(module):
    global _PNG_RED, _PNG_BLUE, _PNG_GREEN
    _PNG_RED = _png_of((200, 30, 30))
    _PNG_BLUE = _png_of((30, 60, 200))
    _PNG_GREEN = _png_of((30, 180, 60))


def _md5(data: bytes) -> str:
    return hashlib.md5(data).hexdigest()


def _md5_file(path) -> str:
    return _md5(open(str(path), "rb").read())


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _make_pptx(path, png1, png2, duplicate_png1=None):
    """slide1: 文本 + group 内 png1(png1 唯一来源在 group 内,递归测试不因
    去重巧合而假绿);slide2: png2 + 可选重复 png1(顶层,供去重用例)。"""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s1.shapes.add_textbox(Emu(0), Emu(0), Emu(1000000), Emu(200000))
    tb.text_frame.text = "第一页说明文字"
    group = s1.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(png1), Emu(0), Emu(0))

    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_picture(io.BytesIO(png2), Emu(200000), Emu(200000))
    if duplicate_png1 is not None:
        s2.shapes.add_picture(io.BytesIO(duplicate_png1), Emu(400000), Emu(200000))
    prs.save(str(path))


class TestPptxExtraction:
    @pytest.fixture()
    def pptx_case(self, tmp_path):
        src = tmp_path / "deck.pptx"
        _make_pptx(src, _PNG_RED, _PNG_BLUE)
        out = tmp_path / "out"
        return src, out, extract_office_images(src, "file_px", out)

    def test_lossless_bytes(self, pptx_case):
        """核心断言:提取产物与 media 原始字节逐位一致(无损证明)。"""
        src, out, images = pptx_case
        assert len(images) == 2  # slide1 group 内 png_red + slide2 png_blue
        assert _md5_file(images[0].image_path) == _md5(_PNG_RED)
        assert _md5_file(images[1].image_path) == _md5(_PNG_BLUE)

    def test_page_numbers_follow_slides(self, pptx_case):
        _, _, images = pptx_case
        assert images[0].page == 1
        assert images[1].page == 2

    def test_slide_text_becomes_context(self, pptx_case):
        _, _, images = pptx_case
        assert "第一页说明文字" in images[0].after_text
        assert images[0].before_text == ""

    def test_dimensions_and_original_ext(self, pptx_case):
        _, _, images = pptx_case
        assert (images[0].width_px, images[0].height_px) == (320, 240)
        assert images[0].image_path.endswith(".png")

    def test_group_shape_picture_is_extracted(self, pptx_case):
        """group 内嵌图同样提取:png_red 唯一来源在 slide1 的 group 内,
        递归失效时该图直接缺失 → md5 断言必然失败(真覆盖,非去重巧合)。"""
        _, _, images = pptx_case
        assert len(images) == 2
        assert images[0].page == 1
        assert _md5_file(images[0].image_path) == _md5(_PNG_RED)

    def test_placeholder_picture_is_extracted(self, tmp_path):
        """版式占位符图(模板化 PPT 极常见):PlaceholderPicture 的 shape_type
        返回 PLACEHOLDER 而非 PICTURE(MRO 中 _InheritsDimensions 覆盖),
        按 shape_type 判断会整类静默漏图;isinstance(Picture) 才是真判据。"""
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
        from pptx.shapes.picture import Picture

        prs = Presentation()
        layout = next(
            l for l in prs.slide_layouts
            if any(ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE
                   for ph in l.placeholders)
        )
        slide = prs.slides.add_slide(layout)
        ph_idx = next(
            ph.placeholder_format.idx for ph in layout.placeholders
            if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        )
        slide.placeholders[ph_idx].insert_picture(io.BytesIO(_PNG_RED))
        src = tmp_path / "ph.pptx"
        prs.save(str(src))

        # 陷阱自证:该 shape isinstance(Picture) 为真但 shape_type 是 PLACEHOLDER
        ph_shape = next(
            sh for sh in slide.shapes
            if isinstance(sh, Picture)
            and sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
        )
        assert ph_shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER

        images = extract_office_images(src, "file_ph", tmp_path / "out")
        assert len(images) == 1
        assert _md5_file(images[0].image_path) == _md5(_PNG_RED)

    def test_corrupt_container_raises(self, tmp_path):
        bad = tmp_path / "broken.pptx"
        bad.write_bytes(b"not a zip container")
        with pytest.raises(Exception):
            extract_office_images(bad, "file_bad", tmp_path / "out")

    def test_unsupported_ext_rejected(self, tmp_path):
        f = tmp_path / "sheet.xlsx"
        f.write_bytes(b"PK\x03\x04")
        with pytest.raises(ValueError):
            extract_office_images(f, "file_x", tmp_path / "out")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _make_docx(path, png1, png2, png3_rid_blob=None):
    """段落图(png1,inline)+ 前后段落 + 表格内图(png2)+ 浮动图(png3,anchor)。"""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsmap, qn
    from docx.shared import Emu

    doc = Document()
    doc.add_paragraph("图前段落一")
    doc.add_paragraph("图前段落二")
    p_img = doc.add_paragraph()
    p_img.add_run().add_picture(io.BytesIO(png1), width=Emu(914400))
    doc.add_paragraph("图后段落一")
    doc.add_paragraph("图后段落二")
    doc.add_paragraph("图后段落三")

    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(png2))

    rid3 = None
    if png3_rid_blob:
        rid3, _ = doc.part.get_or_add_image(io.BytesIO(png3_rid_blob))
        anchor = parse_xml(
            '<w:drawing xmlns:w="%(w)s" xmlns:wp="%(wp)s" xmlns:a="%(a)s" xmlns:r="%(r)s" xmlns:pic="%(pic)s">'
            "<wp:anchor distT=\"0\" distB=\"0\" distL=\"0\" distR=\"0\" simplePos=\"0\" "
            'relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            "<wp:simplePos x=\"0\" y=\"0\"/>"
            '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
            "<a:graphic><a:graphicData>"
            "<pic:pic><a:blip r:embed=\"%(rid)s\"/></pic:pic>"
            "</a:graphicData></a:graphic>"
            "</wp:anchor></w:drawing>"
            % {
                "w": nsmap["w"], "wp": nsmap["wp"], "a": nsmap["a"],
                "r": nsmap["r"], "pic": nsmap["pic"], "rid": rid3,
            }
        )
        doc.add_paragraph()._p.append(anchor)

    doc.save(str(path))


class TestDocxExtraction:
    @pytest.fixture()
    def docx_case(self, tmp_path):
        src = tmp_path / "doc.docx"
        _make_docx(src, _PNG_RED, _PNG_BLUE, png3_rid_blob=_PNG_GREEN)
        out = tmp_path / "out"
        return src, out, extract_office_images(src, "file_dx", out)

    def test_lossless_bytes_all_sources(self, docx_case):
        """inline、表格内、浮动(anchor)三路图全部无损提取。"""
        _, _, images = docx_case
        assert len(images) == 3
        md5s = {_md5_file(i.image_path) for i in images}
        assert md5s == {_md5(_PNG_RED), _md5(_PNG_BLUE), _md5(_PNG_GREEN)}

    def test_page_is_none_for_flow_document(self, docx_case):
        _, _, images = docx_case
        assert all(i.page is None for i in images)

    def test_paragraph_context_windows(self, docx_case):
        _, _, images = docx_case
        first = next(i for i in images if _md5_file(i.image_path) == _md5(_PNG_RED))
        assert "图前段落一" in first.before_text and "图前段落二" in first.before_text
        assert "图后段落一" in first.after_text
        # 上下文窗口 = 各 _CONTEXT_BLOCKS 个非空段落,不越界取满全部
        assert "图后段落三" not in first.after_text

    def test_context_block_limit_is_two(self):
        assert _CONTEXT_BLOCKS == 2

    def test_corrupt_container_raises(self, tmp_path):
        from docx.opc.exceptions import PackageNotFoundError

        bad = tmp_path / "broken.docx"
        bad.write_bytes(b"not a zip container")
        with pytest.raises(PackageNotFoundError):
            extract_office_images(bad, "file_bad", tmp_path / "out")


# ---------------------------------------------------------------------------
# 通用行为
# ---------------------------------------------------------------------------

class TestCommonBehavior:
    def test_unreadable_image_skipped_without_residue(self, tmp_path):
        """PIL 不可识别字节(WMF/EMF/损坏)跳过且不残留写盘文件。"""
        out = tmp_path / "out"
        img_path, w, h = _save_and_measure(b"\xd7\xcd\xc6\xce", "wmf", 3, 0, "file_u", out)
        assert img_path is None and w is None and h is None
        assert not (out / "page_3_img_0.wmf").exists()

    def test_dedup_across_same_blob(self, tmp_path):
        """同字节图多处引用只登记首次:png_red 在 slide1 group 内与 slide2
        顶层各出现一次,仅登记首次;image_id 连续编号无空洞。"""
        src = tmp_path / "deck.pptx"
        _make_pptx(src, _PNG_RED, _PNG_BLUE, duplicate_png1=_PNG_RED)
        images = extract_office_images(src, "file_dd", tmp_path / "out")
        assert [i.image_id for i in images] == ["img_file_dd_0", "img_file_dd_1"]
        assert [_md5_file(i.image_path) for i in images] == [
            _md5(_PNG_RED), _md5(_PNG_BLUE),
        ]

    def test_orphan_files_reconciled_on_success(self, tmp_path):
        """成功路径对账:上一轮中断遗留的孤儿(已写盘未登记,如首提中断后
        retry 提取数收缩)被清理;本轮登记文件保留。"""
        out = tmp_path / "out"
        out.mkdir()
        (out / "page_9_img_9.png").write_bytes(b"orphan from interrupted run")

        src = tmp_path / "deck.pptx"
        _make_pptx(src, _PNG_RED, _PNG_BLUE)
        images = extract_office_images(src, "file_or", out)

        assert not (out / "page_9_img_9.png").exists()
        assert all(Path(i.image_path).exists() for i in images)

    def test_no_reconcile_on_failed_extraction(self, tmp_path):
        """失败路径保守:容器级异常时不动已存在文件,留待下次成功对账。"""
        out = tmp_path / "out"
        out.mkdir()
        (out / "page_9_img_9.png").write_bytes(b"orphan")
        bad = tmp_path / "broken.pptx"
        bad.write_bytes(b"not a zip container")
        with pytest.raises(Exception):
            extract_office_images(bad, "file_or2", out)
        assert (out / "page_9_img_9.png").exists()


# ---------------------------------------------------------------------------
# doc_parser 接线分支(office 图片路 / 文字路分发、失败声明)
# ---------------------------------------------------------------------------

from types import SimpleNamespace

import knowledge.doc_parser as doc_parser_module
from knowledge.doc_parser import DocumentParser, ImageContext


class TestDocParserWiring:
    def _dispatch_parser(self):
        """双 stub converter 的裸 parser(不触发 docling 构建),
        返回 (parser, office_called, pic_called) 供分发断言。"""
        parser = DocumentParser.__new__(DocumentParser)
        office_called, pic_called = [], []

        def _office_convert(p):
            office_called.append(p)
            return SimpleNamespace(document=object())

        def _pic_convert(p):
            pic_called.append(p)
            return SimpleNamespace(document=object())

        parser._office_converter = SimpleNamespace(convert=_office_convert)
        parser.converter = SimpleNamespace(convert=_pic_convert)
        parser._chunker = None
        parser._tokenizer = None
        return parser, office_called, pic_called

    def _parse_docx(self, monkeypatch, tmp_path, extract_side_effect):
        """跑 parse_file(DOCX,降级直喂),返回 (result, extract_calls)。"""
        src = tmp_path / "doc.docx"
        src.write_bytes(b"PK\x03\x04")
        parser, _, _ = self._dispatch_parser()
        monkeypatch.setattr(
            parser, "_convert_office_to_pdf", lambda p: (None, None)
        )  # 降级:无 Office 转换器,直喂原文件
        monkeypatch.setattr(parser, "_extract_chunks", lambda doc: [])
        extract_calls = []

        def _fake_extract(path, file_id, out_dir):
            extract_calls.append((path, file_id, out_dir))
            if isinstance(extract_side_effect, Exception):
                raise extract_side_effect
            return extract_side_effect

        monkeypatch.setattr(doc_parser_module, "extract_office_images", _fake_extract)
        result = parser.parse_file(str(src), "file_w")
        return result, extract_calls

    def test_docx_images_come_from_native_extractor(self, monkeypatch, tmp_path):
        ctx = ImageContext(
            image_id="img_f_0", file_id="f", page=1,
            image_path="x.png", caption=None, before_text="", after_text="",
        )
        result, extract_calls = self._parse_docx(monkeypatch, tmp_path, [ctx])
        assert result.images == [ctx]
        warnings = result.metadata.get("warnings") or []
        assert not any("图片" in w for w in warnings)  # 成功提取不产生图片类 warning
        # 原生提取输入 = 源文件(非转换产物),输出目录 = 同目录
        assert extract_calls[0][1] == "file_w"
        assert str(extract_calls[0][0]).endswith("doc.docx")

    def test_extract_failure_yields_zero_images_with_warning(self, monkeypatch, tmp_path):
        result, _ = self._parse_docx(
            monkeypatch, tmp_path, RuntimeError("zip bomb")
        )
        assert result.images == []
        warnings = result.metadata.get("warnings") or []
        assert any("图片无损提取失败" in w for w in warnings)

    def test_pdf_images_still_from_layout_crop(self, monkeypatch, tmp_path):
        """非 Office 输入维持版面裁剪提取,不走原生提取器。"""
        src = tmp_path / "doc.pdf"
        src.write_bytes(b"%PDF-1.4")
        parser, _, _ = self._dispatch_parser()
        monkeypatch.setattr(parser, "_extract_chunks", lambda doc: [])
        cropped = [ImageContext(
            image_id="img_f_0", file_id="f", page=1,
            image_path="x.png", caption=None, before_text="", after_text="",
        )]
        monkeypatch.setattr(parser, "_extract_images", lambda doc, fid, out: cropped)

        def _fail_extract(*a, **k):
            raise AssertionError("native extractor must not run for PDF")

        monkeypatch.setattr(doc_parser_module, "extract_office_images", _fail_extract)
        result = parser.parse_file(str(src), "file_pdf")
        assert result.images == cropped

    def test_office_routes_to_no_picture_converter_after_pdf_conversion(
        self, monkeypatch, tmp_path
    ):
        """高-1 回归锁:Office 成功转 PDF 后仍必须走无图单例——分发依据
        源格式而非转换产物后缀(file_path 已是 temp_pdf,按后缀分发会让
        带图单例在主路径永远脱靶)。集成方式跑 parse_file,不 mock 分发点。"""
        parser, office_called, pic_called = self._dispatch_parser()
        temp_pdf = tmp_path / "deck.pdf"
        temp_pdf.write_bytes(b"%PDF-1.4")
        monkeypatch.setattr(
            parser, "_convert_office_to_pdf", lambda p: (temp_pdf, None)
        )
        monkeypatch.setattr(parser, "_extract_chunks", lambda doc: [])
        monkeypatch.setattr(parser, "_extract_images", lambda doc, fid, out: [])
        monkeypatch.setattr(
            doc_parser_module, "extract_office_images", lambda *a, **k: []
        )

        src = tmp_path / "deck.pptx"
        src.write_bytes(b"PK\x03\x04")
        parser.parse_file(str(src), "file_o")

        assert len(office_called) == 1
        assert str(office_called[0]).endswith("deck.pdf")  # 产物 PDF 而非源文件
        assert pic_called == []

    def test_non_office_routes_to_picture_converter(self, monkeypatch, tmp_path):
        """PDF 输入走带图单例,原生提取器不得启动。"""
        parser, office_called, pic_called = self._dispatch_parser()
        monkeypatch.setattr(parser, "_extract_chunks", lambda doc: [])
        cropped = [ImageContext(
            image_id="img_f_0", file_id="f", page=1,
            image_path="x.png", caption=None, before_text="", after_text="",
        )]
        monkeypatch.setattr(parser, "_extract_images", lambda doc, fid, out: cropped)

        def _fail_extract(*a, **k):
            raise AssertionError("native extractor must not run for PDF")

        monkeypatch.setattr(doc_parser_module, "extract_office_images", _fail_extract)
        src = tmp_path / "doc.pdf"
        src.write_bytes(b"%PDF-1.4")
        result = parser.parse_file(str(src), "file_pdf")

        assert result.images == cropped
        assert len(pic_called) == 1 and str(pic_called[0]).endswith("doc.pdf")
        assert office_called == []

    def test_office_suffix_match_is_case_insensitive(self, monkeypatch, tmp_path):
        """大写 .PPTX 扩展名同样按 Office 分发(源格式判断 lower())。"""
        parser, office_called, pic_called = self._dispatch_parser()
        monkeypatch.setattr(parser, "_convert_office_to_pdf", lambda p: (None, None))
        monkeypatch.setattr(parser, "_extract_chunks", lambda doc: [])
        monkeypatch.setattr(parser, "_extract_images", lambda doc, fid, out: [])
        monkeypatch.setattr(
            doc_parser_module, "extract_office_images", lambda *a, **k: []
        )

        src = tmp_path / "DECK.PPTX"
        src.write_bytes(b"PK\x03\x04")
        parser.parse_file(str(src), "file_up")

        assert len(office_called) == 1  # 降级直喂源文件
        assert str(office_called[0]).endswith("DECK.PPTX")
        assert pic_called == []
