"""python/tests/test_renderer.py — Markdown→docx 渲染单测（§3.7）。

回归覆盖：
  - B5：_add_text_with_refs 处理 [n] 引用时曾缺失 Pt 导入触发 NameError，
    现函数内自持导入，应正常生成上标 run。
"""
import os
import sys
import tempfile

# Ensure python/ is on sys.path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document  # noqa: E402

from report_generator.renderer import (  # noqa: E402
    _add_text_with_refs,
    render_markdown_to_docx,
)


class TestAddTextWithRefs:
    def test_ref_superscript_no_nameerror(self):
        """含 [n] 引用的文本应拆分为上标 run（回归 B5 NameError）。"""
        doc = Document()
        p = doc.add_paragraph()
        _add_text_with_refs(p, "结论[1]与[2]。")
        texts = [r.text for r in p.runs]
        assert "[1]" in texts
        assert "[2]" in texts
        # 上标 run 应设置 superscript
        sup_runs = [r for r in p.runs if r.font.superscript]
        assert len(sup_runs) == 2

    def test_plain_text_no_ref(self):
        """无引用文本原样输出，不产生上标 run。"""
        doc = Document()
        p = doc.add_paragraph()
        _add_text_with_refs(p, "纯文本无引用")
        assert "".join(r.text for r in p.runs) == "纯文本无引用"
        assert not any(r.font.superscript for r in p.runs)


class TestRenderMarkdownToDocx:
    def test_render_with_refs_end_to_end(self):
        """端到端渲染含引用的 Markdown，应成功产出非空 docx（回归 B5）。"""
        md = "# 标题\n\n## 章节一\n\n正文内容[1]。\n\n- 列表项[2]\n"
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "report.docx")
            result = render_markdown_to_docx(md, out, title="测试报告")
            assert os.path.exists(result)
            assert os.path.getsize(result) > 0
